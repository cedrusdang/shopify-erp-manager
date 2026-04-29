"""Persistent live database screen for the fixed local Excel file."""

from __future__ import annotations

from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox
import math

import openpyxl

from ..backup import open_file
from ..constants import DATABASE_FILE


class LiveDatabasePanel(ttk.LabelFrame):
    """Always-visible local database viewer with live refresh and inline edit."""

    def __init__(
        self,
        parent: tk.Widget,
        app,
        *,
        title: str = " Live Database Screen (Always On) ",
        max_rows: int = 300,
    ):
        super().__init__(parent, text=title, padding=8)
        self._app = app
        self._db_path = Path(DATABASE_FILE)
        self._poll_ms = 2000
        self._poll_job: str | None = None
        self._last_mtime: float | None = None
        self._max_rows = max_rows
        self._rows_per_page = max(1, max_rows)
        self._current_page = 1
        self._total_data_rows = 0
        self._total_pages = 1
        self._header: list[str] = []
        self._all_rows: list[tuple[int, list[str]]] = []

        self._search_text = tk.StringVar(value="")
        self._search_col = tk.StringVar(value="All Columns")
        self._page_input = tk.IntVar(value=1)
        self._page_info = tk.StringVar(value="Page 1/1")
        self._is_applying_filter = False

        self._edit_entry: ttk.Entry | None = None
        self._edit_info: tuple[str, int] | None = None

        self._build()

    def _build(self) -> None:
        top = ttk.Frame(self)
        top.pack(fill=tk.X, pady=(0, 6))

        ttk.Label(
            top,
            text="(5) Local database is always visible below. Drag the middle divider to resize.",
            font=("Segoe UI", 9, "bold"),
            foreground="#0f6b45",
        ).pack(side=tk.LEFT)

        ttk.Button(top, text="Refresh Now", command=self.refresh_data).pack(side=tk.RIGHT, padx=(6, 0))
        ttk.Button(top, text="Open in Excel", command=self._open_excel).pack(side=tk.RIGHT)

        self._status = tk.StringVar(value=f"File: {self._db_path.name}")
        ttk.Label(
            self,
            textvariable=self._status,
            foreground="#666",
            font=("Segoe UI", 8),
        ).pack(anchor=tk.W, pady=(0, 4))

        # Search / filter controls
        tools = ttk.Frame(self)
        tools.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(tools, text="Search:").pack(side=tk.LEFT)
        search_entry = ttk.Entry(tools, textvariable=self._search_text, width=30)
        search_entry.pack(side=tk.LEFT, padx=(4, 8))
        # Use trace instead of KeyRelease for more reliable filtering
        self._search_text.trace("w", lambda *_: self._apply_filter())

        ttk.Label(tools, text="Column:").pack(side=tk.LEFT)
        self._col_combo = ttk.Combobox(
            tools,
            textvariable=self._search_col,
            state="readonly",
            width=26,
            values=["All Columns"],
        )
        self._col_combo.pack(side=tk.LEFT, padx=(4, 8))
        # Use trace for more reliable filtering on column changes
        self._search_col.trace("w", lambda *_: self._apply_filter())

        ttk.Button(tools, text="Clear", command=self._clear_filter).pack(side=tk.LEFT)

        ttk.Separator(tools, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=8)
        ttk.Button(tools, text="◀ Prev", command=self._prev_page).pack(side=tk.LEFT)
        ttk.Label(tools, textvariable=self._page_info).pack(side=tk.LEFT, padx=(8, 8))
        ttk.Label(tools, text="Go to:").pack(side=tk.LEFT)
        self._page_spin = ttk.Spinbox(tools, from_=1, to=1, textvariable=self._page_input, width=7)
        self._page_spin.pack(side=tk.LEFT, padx=(4, 4))
        ttk.Button(tools, text="Go", command=self._goto_page).pack(side=tk.LEFT)
        ttk.Button(tools, text="Next ▶", command=self._next_page).pack(side=tk.LEFT, padx=(6, 0))

        grid_wrap = ttk.Frame(self)
        grid_wrap.pack(fill=tk.BOTH, expand=True)

        self._tree = ttk.Treeview(grid_wrap, show="headings")
        self._tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self._tree.bind("<Double-1>", self._begin_edit)

        ysb = ttk.Scrollbar(grid_wrap, orient=tk.VERTICAL, command=self._tree.yview)
        xsb = ttk.Scrollbar(self, orient=tk.HORIZONTAL, command=self._tree.xview)
        self._tree.configure(yscrollcommand=ysb.set, xscrollcommand=xsb.set)
        ysb.pack(side=tk.RIGHT, fill=tk.Y)
        xsb.pack(fill=tk.X)

    def start_auto_refresh(self) -> None:
        self.refresh_data()
        self._schedule_poll()

    def stop_auto_refresh(self) -> None:
        if self._poll_job:
            self.after_cancel(self._poll_job)
            self._poll_job = None

    def _schedule_poll(self) -> None:
        self._poll_job = self.after(self._poll_ms, self._poll)

    def _poll(self) -> None:
        try:
            mtime = self._db_path.stat().st_mtime if self._db_path.exists() else None
        except OSError:
            mtime = None

        if mtime != self._last_mtime:
            self.refresh_data()

        self._schedule_poll()

    def _open_excel(self) -> None:
        if not open_file(self._db_path):
            messagebox.showwarning("Not Found", f"File not found:\n{self._db_path.resolve()}", parent=self)

    def refresh_data(self) -> None:
        if not self._db_path.exists():
            self._clear_tree()
            self._status.set(f"Waiting for {self._db_path.name} (file not found)")
            self._last_mtime = None
            return

        try:
            wb = openpyxl.load_workbook(self._db_path, read_only=True, data_only=True)
            ws = wb.active
            self._total_data_rows = max(0, int(ws.max_row) - 1)

            rows = list(ws.iter_rows(values_only=True, min_row=1, max_row=1))
            self._total_pages = max(1, math.ceil(self._total_data_rows / self._rows_per_page))
            if self._current_page > self._total_pages:
                self._current_page = self._total_pages
            if self._current_page < 1:
                self._current_page = 1

            start_data_row = 2 + (self._current_page - 1) * self._rows_per_page
            end_data_row = min(start_data_row + self._rows_per_page - 1, self._total_data_rows + 1)

            if self._total_data_rows > 0:
                data_rows = list(
                    ws.iter_rows(values_only=True, min_row=start_data_row, max_row=end_data_row)
                )
            else:
                data_rows = []
        except Exception as exc:
            self._status.set(f"Read error: {exc}")
            return

        if not rows:
            self._clear_tree()
            self._status.set(f"{self._db_path.name} is empty")
            return

        header = [str(c or "") for c in rows[0]]
        if not any(header):
            col_count = max(len(r) for r in rows)
            header = [f"col_{i}" for i in range(1, col_count + 1)]

        data = data_rows
        self._set_columns(header)
        self._header = header
        self._all_rows = []
        for i, row_vals in enumerate(data):
            idx = 2 + (self._current_page - 1) * self._rows_per_page + i
            row = ["" if v is None else str(v) for v in row_vals]
            if len(row) < len(header):
                row.extend([""] * (len(header) - len(row)))
            self._all_rows.append((idx, row[: len(header)]))

        col_values = ["All Columns"] + [h or f"col_{i+1}" for i, h in enumerate(header)]
        self._col_combo.configure(values=col_values)
        if self._search_col.get() not in col_values:
            self._search_col.set("All Columns")

        self._page_spin.configure(to=self._total_pages)
        self._page_input.set(self._current_page)
        self._page_info.set(f"Page {self._current_page}/{self._total_pages}")

        self._apply_filter()

        try:
            self._last_mtime = self._db_path.stat().st_mtime
        except OSError:
            self._last_mtime = None

        shown = len(self._tree.get_children())
        start_row = 0 if self._total_data_rows == 0 else (self._current_page - 1) * self._rows_per_page + 1
        end_row = min(self._current_page * self._rows_per_page, self._total_data_rows)
        self._status.set(
            f"Live view: {self._db_path.name} | total rows: {self._total_data_rows} | page {self._current_page}/{self._total_pages} | rows {start_row}-{end_row} | showing {shown}"
        )

    def _clear_filter(self) -> None:
        self._search_text.set("")
        self._search_col.set("All Columns")
        self._apply_filter()

    def _apply_filter(self) -> None:
        if self._is_applying_filter:
            return
        self._is_applying_filter = True
        try:
            query = self._search_text.get().strip().lower()
            selected_col = self._search_col.get().strip() or "All Columns"

            if selected_col == "All Columns":
                col_index = None
            else:
                col_index = None
                for i, h in enumerate(self._header):
                    name = h or f"col_{i+1}"
                    if name == selected_col:
                        col_index = i
                        break

            for item in list(self._tree.get_children()):
                try:
                    self._tree.delete(item)
                except tk.TclError:
                    # Tree items can change during refresh/trace callbacks; skip stale ids.
                    pass

            matched = 0
            for excel_row, row in self._all_rows:
                if not query:
                    keep = True
                else:
                    if col_index is None:
                        keep = any(query in str(v).lower() for v in row)
                    else:
                        keep = col_index < len(row) and query in str(row[col_index]).lower()

                if keep:
                    self._tree.insert("", tk.END, iid=str(excel_row), values=row)
                    matched += 1

            self._status.set(
                f"Live view: {self._db_path.name} | total rows: {self._total_data_rows} | page {self._current_page}/{self._total_pages} | filtered rows shown: {matched}"
            )
        finally:
            self._is_applying_filter = False

    def _prev_page(self) -> None:
        if self._current_page > 1:
            self._current_page -= 1
            self.refresh_data()

    def _next_page(self) -> None:
        if self._current_page < self._total_pages:
            self._current_page += 1
            self.refresh_data()

    def _goto_page(self) -> None:
        page = max(1, min(self._total_pages, int(self._page_input.get() or 1)))
        self._current_page = page
        self.refresh_data()

    def _clear_tree(self) -> None:
        self._tree["columns"] = ()
        for item in self._tree.get_children():
            self._tree.delete(item)

    def _set_columns(self, header: list[str]) -> None:
        columns = [f"c{i}" for i in range(1, len(header) + 1)]
        self._tree["columns"] = columns
        for idx, col in enumerate(columns):
            title = header[idx] or f"col_{idx + 1}"
            self._tree.heading(col, text=title)
            self._tree.column(col, width=130, anchor=tk.W, stretch=True)

    def _begin_edit(self, event: tk.Event) -> None:
        row_id = self._tree.identify_row(event.y)
        col_id = self._tree.identify_column(event.x)
        if not row_id or not col_id:
            return

        col_index = int(col_id.replace("#", ""))
        bbox = self._tree.bbox(row_id, col_id)
        if not bbox:
            return
        x, y, w, h = bbox

        old_values = list(self._tree.item(row_id, "values"))
        old_val = old_values[col_index - 1] if col_index - 1 < len(old_values) else ""

        if self._edit_entry:
            self._edit_entry.destroy()
            self._edit_entry = None

        self._edit_entry = ttk.Entry(self._tree)
        self._edit_entry.insert(0, old_val)
        self._edit_entry.place(x=x, y=y, width=w, height=h)
        self._edit_entry.focus_set()
        self._edit_info = (row_id, col_index)

        self._edit_entry.bind("<Return>", lambda _: self._commit_edit())
        self._edit_entry.bind("<Escape>", lambda _: self._cancel_edit())
        self._edit_entry.bind("<FocusOut>", lambda _: self._commit_edit())

    def _cancel_edit(self) -> None:
        if self._edit_entry:
            self._edit_entry.destroy()
            self._edit_entry = None
        self._edit_info = None

    def _commit_edit(self) -> None:
        if not self._edit_entry or not self._edit_info:
            return

        row_id, col_index = self._edit_info
        new_val = self._edit_entry.get()
        self._edit_entry.destroy()
        self._edit_entry = None
        self._edit_info = None

        values = list(self._tree.item(row_id, "values"))
        while len(values) < col_index:
            values.append("")
        values[col_index - 1] = new_val
        self._tree.item(row_id, values=values)

        excel_row = int(row_id)
        if not self._write_cell(excel_row, col_index, new_val):
            self.refresh_data()

    def _write_cell(self, row: int, col: int, value: str) -> bool:
        try:
            wb = openpyxl.load_workbook(self._db_path)
            ws = wb.active
            ws.cell(row=row, column=col, value=value)
            wb.save(self._db_path)
            wb.close()
            try:
                self._last_mtime = self._db_path.stat().st_mtime
            except OSError:
                pass
            self._app.set_status(f"Updated local DB cell R{row}C{col}")
            return True
        except Exception as exc:
            messagebox.showerror(
                "Write Error",
                f"Could not write to {self._db_path.name}:\n{exc}",
                parent=self,
            )
            return False
