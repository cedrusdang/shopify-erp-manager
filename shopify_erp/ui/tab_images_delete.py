"""Image delete tools for Shopify product gallery fields and product metafields."""

from __future__ import annotations

import logging
import re
import threading
import time
from tkinter import scrolledtext
import tkinter as tk
from tkinter import ttk, messagebox

import openpyxl

from ..api import clear_product_metafield, delete_product_image_by_field
from ..constants import DATABASE_FILE, DEFAULT_SKU

logger = logging.getLogger(__name__)


class ImagesDeleteTab(ttk.Frame):
    """Delete Shopify images by selected product gallery field or product metafield."""

    def __init__(self, parent: ttk.Notebook, app):
        super().__init__(parent)
        self._app = app
        self._running = False
        self._field_pick = tk.StringVar(value="")
        self._sku_mode = tk.StringVar(value="all")
        self._sku_pick = tk.StringVar(value="")
        self._sku_group = tk.StringVar(value="")
        self._stats = tk.StringVar(value="")
        self._build()

    def _build(self) -> None:
        ttk.Label(
            self,
            text="Image Delete — delete Shopify gallery images or clear product image metafields",
            font=("Segoe UI", 10, "bold"),
            foreground="#8a1f11",
        ).pack(anchor=tk.W, padx=10, pady=(10, 6))

        info = ttk.LabelFrame(self, text=" How It Works ", padding=8)
        info.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(
            info,
            text=(
                "1. Select a target field from the database headers.\n"
                "2. images means clear the current product gallery; images.0.src means delete the first gallery slot only.\n"
                "3. metafields.namespace.key means delete that product metafield entry.\n"
                "4. SKU scope can be All, Single, or Group.\n"
                "5. This tab does not delete variant metafields or Shopify Files records."
            ),
            justify=tk.LEFT,
            foreground="#444",
            font=("Segoe UI", 9),
        ).pack(anchor=tk.W)

        options = ttk.LabelFrame(self, text=" Delete Options ", padding=8)
        options.pack(fill=tk.X, padx=8, pady=(4, 2))

        field_row = ttk.Frame(options)
        field_row.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(field_row, text="Target Field:").pack(side=tk.LEFT)
        self._field_combo = ttk.Combobox(field_row, textvariable=self._field_pick, state="readonly", width=42)
        self._field_combo.pack(side=tk.LEFT, padx=(6, 0), fill=tk.X, expand=True)
        ttk.Button(field_row, text="Refresh Fields", command=self._load_field_options).pack(side=tk.LEFT, padx=(6, 0))

        ttk.Label(
            options,
            text=(
                "Supported delete targets: product gallery fields (images, images.N.src) and product metafields "
                "(metafields.namespace.key)."
            ),
            foreground="#555",
            font=("Segoe UI", 8),
            justify=tk.LEFT,
        ).pack(anchor=tk.W, pady=(0, 6))

        sku_row = ttk.Frame(options)
        sku_row.pack(fill=tk.X, pady=(0, 2))
        ttk.Label(sku_row, text="SKU Scope:").pack(side=tk.LEFT)
        ttk.Radiobutton(sku_row, text="All", value="all", variable=self._sku_mode, command=self._sync_sku_mode).pack(side=tk.LEFT, padx=(6, 8))
        ttk.Radiobutton(sku_row, text="Single", value="single", variable=self._sku_mode, command=self._sync_sku_mode).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Radiobutton(sku_row, text="Group", value="group", variable=self._sku_mode, command=self._sync_sku_mode).pack(side=tk.LEFT, padx=(0, 8))

        single_sku_row = ttk.Frame(options)
        single_sku_row.pack(fill=tk.X, pady=(0, 2))
        ttk.Label(single_sku_row, text="Single SKU:").pack(side=tk.LEFT)
        self._sku_combo = ttk.Combobox(single_sku_row, textvariable=self._sku_pick, state="disabled", width=36)
        self._sku_combo.pack(side=tk.LEFT, padx=(6, 0), fill=tk.X, expand=True)

        group_sku_row = ttk.Frame(options)
        group_sku_row.pack(fill=tk.X, pady=(0, 4))
        ttk.Label(group_sku_row, text="Group SKU (comma/newline):").pack(side=tk.LEFT)
        self._sku_group_entry = ttk.Entry(group_sku_row, textvariable=self._sku_group, state="disabled")
        self._sku_group_entry.pack(side=tk.LEFT, padx=(6, 0), fill=tk.X, expand=True)

        btns = ttk.Frame(self)
        btns.pack(fill=tk.X, padx=8, pady=4)
        ttk.Button(
            btns,
            text="Delete Images",
            style="Primary.TButton",
            command=self._start_delete_images,
        ).pack(side=tk.LEFT, padx=(0, 8))

        ttk.Label(self, textvariable=self._stats, foreground="#333").pack(anchor=tk.W, padx=10, pady=(0, 4))

        self._prog = ttk.Progressbar(self, mode="determinate")
        self._prog.pack(fill=tk.X, padx=8, pady=(0, 6))

        ttk.Label(self, text="Delete Log:", font=("Segoe UI", 9, "bold")).pack(anchor=tk.W, padx=8)
        self._log_box = scrolledtext.ScrolledText(self, state="disabled", font=("Consolas", 8))
        self._log_box.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        self._sync_sku_mode()
        self._load_field_options()
        self._load_skus_from_db()

    def _log(self, msg: str) -> None:
        logger.info(f"[ImagesDeleteTab] {msg}")
        try:
            self.after(0, lambda m=msg: self._app.log(self._log_box, m))
        except Exception:
            logger.debug("[ImagesDeleteTab] UI log dispatch skipped", exc_info=True)

    def _ui_async(self, fn) -> None:
        try:
            self.after(0, fn)
        except Exception:
            logger.debug("[ImagesDeleteTab] UI async dispatch failed", exc_info=True)

    def _sync_sku_mode(self) -> None:
        mode = self._sku_mode.get()
        if mode == "single":
            self._sku_combo.config(state="readonly")
            self._sku_group_entry.config(state="disabled")
        elif mode == "group":
            self._sku_combo.config(state="disabled")
            self._sku_group_entry.config(state="normal")
        else:
            self._sku_combo.config(state="disabled")
            self._sku_group_entry.config(state="disabled")

    def _is_product_metafield_field(self, field_name: str) -> bool:
        return field_name.startswith("metafields.")

    def _is_product_image_field(self, field_name: str) -> bool:
        return field_name == "images" or (field_name.startswith("images.") and field_name.endswith(".src"))

    def _is_supported_field(self, field_name: str) -> bool:
        return self._is_product_image_field(field_name) or self._is_product_metafield_field(field_name)

    def _field_options_from_headers(self, headers: list[str]) -> list[str]:
        fields: list[str] = []
        for header in headers:
            key = str(header or "").strip()
            if not key:
                continue
            if self._is_supported_field(key):
                fields.append(key)
        return fields

    def _load_field_options(self) -> None:
        try:
            wb = openpyxl.load_workbook(DATABASE_FILE, read_only=True, data_only=True)
            ws = wb.active
            rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
            first = next(rows, None)
            headers = [str(h or "").strip() for h in (first or [])]
            wb.close()
        except Exception as exc:
            logger.exception("Failed to load delete field options")
            self._log(f"Error loading fields: {exc}")
            return

        fields = self._field_options_from_headers(headers)
        current = self._field_pick.get().strip()
        self._field_combo["values"] = fields
        if current in fields:
            self._field_pick.set(current)
        elif fields:
            self._field_pick.set(fields[0])
        else:
            self._field_pick.set("")

    def _load_skus_from_db(self) -> None:
        wb = None
        try:
            wb = openpyxl.load_workbook(DATABASE_FILE, data_only=True)
            ws = wb.active
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            headers = [str(h or "").strip() for h in (header_row or [])]
            sku_idx = self._find_col_idx(headers, {"variants.0.sku", "sku"})
            if sku_idx is None:
                self._sku_combo["values"] = []
                return

            skus: list[str] = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                sku = str(row[sku_idx]).strip() if len(row) > sku_idx and row[sku_idx] else ""
                if sku and sku != DEFAULT_SKU:
                    skus.append(sku)
            self._sku_combo["values"] = skus
        except Exception as exc:
            logger.exception("Failed to load SKUs from DB for delete tab")
            self._log(f"Error loading SKUs: {exc}")
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

    def _find_col_idx(self, headers: list[str], accepted: set[str]) -> int | None:
        lowered = [h.lower() for h in headers]
        for name in accepted:
            key = name.lower()
            if key in lowered:
                return lowered.index(key)
        return None

    def _selected_sku_set(self) -> set[str]:
        mode = self._sku_mode.get()
        if mode == "single":
            sku = self._sku_pick.get().strip()
            return {sku} if sku else set()
        if mode == "group":
            raw = self._sku_group.get().strip()
            if not raw:
                return set()
            return {item.strip() for item in re.split(r"[,\n]+", raw) if item.strip()}
        return set()

    def _start_delete_images(self) -> None:
        field_name = self._field_pick.get().strip()
        if not field_name:
            messagebox.showwarning("No Field", "Please select target field first.", parent=self)
            return
        if not self._is_supported_field(field_name):
            messagebox.showerror(
                "Unsupported Field",
                "This tab supports only product gallery fields (images, images.N.src) and product metafields (metafields.namespace.key).",
                parent=self,
            )
            return

        store, token = self._app.get_conn()
        if not store or not token:
            return

        sku_set = self._selected_sku_set()
        scope_desc = f"{len(sku_set)} SKU(s)" if sku_set else "all SKUs"
        if field_name == "images":
            delete_mode = "main/product gallery images"
        elif self._is_product_image_field(field_name):
            delete_mode = "product gallery slot"
        else:
            delete_mode = "product metafield"
        if not self._app.confirm_danger(
            "Confirm Image Delete",
            (
                f"Delete target: {field_name}\n"
                f"Delete mode: {delete_mode}\n"
                f"Scope: {scope_desc}\n\n"
                "This cannot be undone from the app. Continue?"
            ),
        ):
            return

        self._running = True
        self._app.start_prog("determinate", use_busy_cursor=False)
        self._prog["maximum"] = 100
        self._prog["value"] = 0
        self._log("Delete task started…")

        def task() -> None:
            try:
                self._run_delete_task(store, token, field_name, sku_set)
            except Exception as exc:
                logger.exception("Delete task crashed")
                self._ui_async(lambda: self._app.stop_prog(clear_busy_cursor=False))
                self._ui_async(lambda: messagebox.showerror("Delete Error", str(exc), parent=self))
            finally:
                self._running = False

        threading.Thread(target=task, daemon=True).start()

    def _run_delete_task(self, store: str, token: str, field_name: str, sku_set: set[str]) -> None:
        start_time = time.time()
        wb = None

        try:
            self._log_step(1, "Loading product database")
            wb = openpyxl.load_workbook(DATABASE_FILE, data_only=True)
            ws = wb.active

            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            headers = [str(h or "").strip() for h in (header_row or [])]
            id_idx = self._find_col_idx(headers, {"id"})
            sku_idx = self._find_col_idx(headers, {"variants.0.sku", "sku"})
            if id_idx is None or sku_idx is None:
                raise RuntimeError("Database must contain 'id' and 'variants.0.sku' (or 'sku') columns.")

            sku_to_id: dict[str, str] = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(row) <= max(id_idx, sku_idx):
                    continue
                prod_id = str(row[id_idx]).strip() if row[id_idx] else ""
                sku = str(row[sku_idx]).strip() if row[sku_idx] else ""
                if prod_id and sku and sku != DEFAULT_SKU:
                    sku_to_id[sku] = prod_id

            target_skus = set(sku_to_id.keys()) if not sku_set else (sku_set & set(sku_to_id.keys()))
            if not target_skus:
                self._ui_async(lambda: self._log("No matching SKUs found in database."))
                self._ui_async(lambda: self._app.stop_prog(clear_busy_cursor=False))
                return

            self._log_step(2, f"Deleting for {len(target_skus)} product(s)")
            ok_count = 0
            fail_count = 0

            for index, sku in enumerate(sorted(target_skus), start=1):
                if not self._running:
                    self._log("Delete cancelled by user")
                    break

                product_id = sku_to_id[sku]
                if self._is_product_image_field(field_name):
                    ok, _status, msg = delete_product_image_by_field(store, token, product_id, field_name)
                else:
                    ok, _status, msg = clear_product_metafield(store, token, product_id, field_name)

                if ok:
                    ok_count += 1
                    self._log(f"  [{index}/{len(target_skus)}] SKU {sku}: cleared ✓")
                else:
                    fail_count += 1
                    self._log(f"  [{index}/{len(target_skus)}] SKU {sku}: failed ✗ ({msg})")

                progress = int(index / len(target_skus) * 100)
                self._ui_async(lambda p=progress: self._set_progress(p, ok_count, fail_count, len(target_skus)))

            elapsed = time.time() - start_time
            self._log_step(3, f"Delete completed in {elapsed:.1f}s (ok={ok_count}, fail={fail_count})")
            self._ui_async(lambda: self._app.stop_prog(clear_busy_cursor=False))
            self._ui_async(
                lambda: messagebox.showinfo(
                    "Delete Complete",
                    f"Deleted/Cleared: {ok_count}\nFailed: {fail_count}\n\nTime: {elapsed:.1f}s",
                    parent=self,
                )
            )
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

    def _set_progress(self, progress: int, ok_count: int, fail_count: int, total: int) -> None:
        self._prog["maximum"] = 100
        self._prog["value"] = progress
        self._app.set_prog_value(progress, 100)
        self._stats.set(f"Processed: {ok_count + fail_count}/{total} | Cleared: {ok_count} | Failed: {fail_count}")

    def _log_step(self, step_no: int, title: str, detail: str = "") -> None:
        if detail:
            self._log(f"[Step {step_no}] {title} — {detail}")
        else:
            self._log(f"[Step {step_no}] {title}")