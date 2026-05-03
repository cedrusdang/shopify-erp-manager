"""
Upload tab — push Excel changes back to Shopify.
"""

from __future__ import annotations

import base64
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
import json
import re
import threading
import time
from pathlib import Path
from tkinter import scrolledtext
import tkinter as tk
from tkinter import ttk, messagebox

import openpyxl

from ..api import update_product_api
from ..constants import DATABASE_FILE, DEFAULT_SKU, IMAGE_DIR, REQUIRED_FIELDS
from ..session import save_session, load_session, clear_session


class UploadTab(ttk.Frame):
    def __init__(self, parent: ttk.Notebook, app):
        super().__init__(parent)
        self._app     = app
        self._running = False
        self._upload_col_vars: dict[str, tk.BooleanVar] = {}
        self._upload_headers: list[str] = []
        self._upload_cols_file_sig: tuple[int, int] | None = None
        self._upload_cols_poll_job: str | None = None
        self._safe_delay = tk.DoubleVar(value=0.3)
        self._pause_event = threading.Event()
        self._stop_event = threading.Event()
        self._paused = False
        self._build()

    def _required_upload_fields(self) -> set[str]:
        """Fields that remain selected/locked in Upload selector.

        Upload only needs product `id` (to identify the row) and
        `variants.0.id` (so variant PUT knows which variant to update).
        Deliberately does NOT include title/handle so that selecting only
        variant columns results in exactly 1 variant request per SKU.
        """
        return {"id", "variants.0.id"}

    # ──────────────────────────────────────────────────
    def _build(self) -> None:
        ttk.Label(
            self,
            text="(4) Upload fixed database file to Shopify",
            font=("Segoe UI", 10, "bold"),
            foreground="#0f6b45",
        ).pack(anchor=tk.W, padx=10, pady=(8, 0))

        # Fixed file display
        frow = ttk.LabelFrame(self, text=" Source Excel File (fixed) ", padding=8)
        frow.pack(fill=tk.X, padx=8, pady=8)
        self._up_path = tk.StringVar(value=DATABASE_FILE)
        ttk.Label(
            frow, textvariable=self._up_path,
            foreground="#555", wraplength=620,
        ).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Label(
            frow,
            text="No browsing allowed",
            foreground="#888",
            font=("Segoe UI", 8, "italic"),
        ).pack(side=tk.RIGHT)

        # Upload field selection
        selector = ttk.LabelFrame(self, text=" Upload Only Selected Fields ", padding=8)
        selector.pack(fill=tk.BOTH, padx=8, pady=(0, 6))
        selector_top = ttk.Frame(selector)
        selector_top.pack(fill=tk.X, pady=(0, 4))
        ttk.Label(
            selector_top,
            text="Choose columns to upload. Required fields are locked and always included.",
            foreground="#555",
            font=("Segoe UI", 8),
        ).pack(side=tk.LEFT)
        ttk.Button(selector_top, text="Load Columns", command=self._load_upload_columns).pack(side=tk.RIGHT, padx=(4, 0))
        ttk.Button(selector_top, text="Clear All", command=self._clear_upload_cols).pack(side=tk.RIGHT, padx=(4, 0))
        ttk.Button(selector_top, text="Select All", command=self._select_all_upload_cols).pack(side=tk.RIGHT)

        self._upload_col_canvas = tk.Canvas(selector, borderwidth=0, highlightthickness=0, height=122)
        upload_col_vsb = ttk.Scrollbar(selector, orient="vertical", command=self._upload_col_canvas.yview)
        self._upload_col_canvas.configure(yscrollcommand=upload_col_vsb.set)
        upload_col_vsb.pack(side=tk.RIGHT, fill=tk.Y)
        self._upload_col_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self._upload_col_inner = ttk.Frame(self._upload_col_canvas)
        self._upload_col_win = self._upload_col_canvas.create_window((0, 0), window=self._upload_col_inner, anchor="nw")
        self._upload_col_inner.bind(
            "<Configure>",
            lambda _e: self._upload_col_canvas.configure(scrollregion=self._upload_col_canvas.bbox("all")),
        )
        self._upload_col_canvas.bind(
            "<Configure>",
            lambda e: self._upload_col_canvas.itemconfig(self._upload_col_win, width=e.width),
        )

        # Buttons
        btns = ttk.Frame(self)
        btns.pack(fill=tk.X, padx=8, pady=2)
        ttk.Button(
            btns, text="(4b) Upload to Shopify",
            command=self._start_upload,
            style="Primary.TButton",
        ).pack(side=tk.LEFT, padx=4)
        self._cont_btn = ttk.Button(
            btns, text="(4c) Continue Upload",
            command=self._continue_upload, state="disabled",
        )
        self._cont_btn.pack(side=tk.LEFT, padx=4)
        self._pause_btn = ttk.Button(
            btns, text="Pause", command=self._toggle_pause_upload, state="disabled"
        )
        self._pause_btn.pack(side=tk.LEFT, padx=4)
        self._stop_btn = ttk.Button(
            btns, text="Stop", command=self._stop_upload_now, state="disabled"
        )
        self._stop_btn.pack(side=tk.LEFT, padx=4)
        self._reset_run_btn = ttk.Button(
            btns, text="Stop + Reset", command=self._stop_and_reset_upload, state="disabled"
        )
        self._reset_run_btn.pack(side=tk.LEFT, padx=4)
        ttk.Button(
            btns, text="✖  Clear Session",
            command=self._clear_session,
        ).pack(side=tk.LEFT, padx=4)

        perf = ttk.Frame(self)
        perf.pack(fill=tk.X, padx=8, pady=(0, 2))
        ttk.Label(perf, text="Safe Delay (s):", foreground="#555").pack(side=tk.LEFT)
        ttk.Spinbox(
            perf,
            from_=0.0,
            to=2.0,
            increment=0.05,
            width=6,
            textvariable=self._safe_delay,
        ).pack(side=tk.LEFT, padx=(6, 12))
        ttk.Label(
            perf,
            text="Lower = faster (higher risk 429); higher = safer",
            foreground="#777",
            font=("Segoe UI", 8),
        ).pack(side=tk.LEFT)

        # Stats + progress
        self._stats = tk.StringVar(value="")
        self._stats_label = tk.Label(
            self, textvariable=self._stats,
            font=("Segoe UI", 9), foreground="#333",
            anchor="w", justify="left", wraplength=700,
        )
        self._stats_label.pack(anchor=tk.W, padx=12, pady=2)
        self._prog = ttk.Progressbar(self, mode="determinate")
        self._prog.pack(fill=tk.X, padx=8, pady=2)

        # Log
        ttk.Label(
            self, text="Upload Log:", font=("Segoe UI", 9, "bold"),
        ).pack(anchor=tk.W, padx=8)
        self._log_box = scrolledtext.ScrolledText(
            self, state="disabled", font=("Consolas", 8),
        )
        self._log_box.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        self._load_upload_columns()
        self._start_upload_columns_auto_refresh()
        self.bind("<Destroy>", self._on_destroy)

    def _on_destroy(self, _event=None) -> None:
        if self._upload_cols_poll_job is not None:
            try:
                self.after_cancel(self._upload_cols_poll_job)
            except Exception:
                pass
            self._upload_cols_poll_job = None

    def _start_upload_columns_auto_refresh(self) -> None:
        def _poll() -> None:
            file_path = Path(DATABASE_FILE)
            if file_path.exists():
                try:
                    stat = file_path.stat()
                    sig = (int(stat.st_mtime_ns), int(stat.st_size))
                except Exception:
                    sig = None
                if sig is not None and sig != self._upload_cols_file_sig:
                    self._load_upload_columns(silent_missing=True)
                    self._log("Upload columns auto-refreshed from latest file.")
            self._upload_cols_poll_job = self.after(2500, _poll)

        self._upload_cols_poll_job = self.after(1200, _poll)

    # ──────────────────────────────────────────────────
    def check_resume_session(self) -> None:
        sess = load_session()
        if sess:
            self._cont_btn.config(state="normal")
            row   = sess.get("pointer", 0)
            total = sess.get("total", "?")
            self._stats.set(
                f"⚠ Upload incomplete! File: {Path(sess['file']).name}  "
                f"— stopped at row {row + 1} / {total}.  "
                f"Click '(4c) Continue Upload' to resume, or '✖ Clear Session' to discard."
            )
            self._stats_label.config(foreground="#cc0000")
        else:
            self._stats_label.config(foreground="#333")

    # ──────────────────────────────────────────────────
    def _log(self, msg: str) -> None:
        self._app.log(self._log_box, msg)

    def _safe_sku(self, value: str) -> str:
        sku = re.sub(r"[^A-Za-z0-9._-]+", "_", (value or "").strip())
        return (sku[:80] or "product")

    def _field_suffix(self, field_name: str) -> str:
        m = re.match(r"^images\.(\d+)\.src$", field_name)
        if m:
            return f"img_{m.group(1)}"
        clean = re.sub(r"[^A-Za-z0-9._-]+", "_", (field_name or "").strip())
        return clean[:40] or "img"

    def _field_folder_name(self, field_name: str) -> str:
        clean = re.sub(r"[^A-Za-z0-9._-]+", "_", (field_name or "").strip())
        return clean[:80] or "images"

    def _build_images_payload_from_folder(self, sku: str, image_fields: list[str] | None = None) -> list[dict]:
        """Fallback: load local images from field folders and convert to attachments."""
        folder = Path(IMAGE_DIR)
        if not folder.exists():
            return []

        prefix = self._safe_sku(sku)
        files: list[Path] = []
        if image_fields:
            for field in image_fields:
                field_folder = folder / self._field_folder_name(field)
                if not field_folder.exists():
                    continue
                suffix = self._field_suffix(field)
                files.extend([p for p in field_folder.glob(f"{prefix}_{suffix}*.*") if p.is_file()])
        else:
            files = [p for p in folder.rglob(f"{prefix}*.*") if p.is_file()]

        files = sorted(set(files))
        payload: list[dict] = []
        for p in files:
            try:
                encoded = base64.b64encode(p.read_bytes()).decode("ascii")
                payload.append({"attachment": encoded, "filename": p.name})
            except Exception:
                continue
        return payload

    def _set_path_value(self, root: dict, path: str, value: str) -> bool:
        """Set dotted-path values dynamically (e.g. variants.0.price) without field whitelists."""
        tokens = [t for t in path.split(".") if t]
        if not tokens:
            return False

        node: dict | list = root
        for i, token in enumerate(tokens):
            is_last = i == len(tokens) - 1
            next_token = tokens[i + 1] if not is_last else None

            if token.isdigit():
                if not isinstance(node, list):
                    return False
                idx = int(token)
                while len(node) <= idx:
                    if is_last:
                        node.append(None)
                    else:
                        node.append([] if next_token and next_token.isdigit() else {})
                if is_last:
                    node[idx] = value
                    return True
                if not isinstance(node[idx], (dict, list)):
                    node[idx] = [] if next_token and next_token.isdigit() else {}
                node = node[idx]
            else:
                if not isinstance(node, dict):
                    return False
                if is_last:
                    node[token] = value
                    return True
                if token not in node or not isinstance(node[token], (dict, list)):
                    node[token] = [] if next_token and next_token.isdigit() else {}
                node = node[token]

        return False

    def _load_upload_columns(self, silent_missing: bool = False) -> None:
        file_path = Path(DATABASE_FILE)
        if not file_path.exists():
            if not silent_missing:
                self._log(f"Upload columns: waiting for {file_path.name}")
            return

        wb = None
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True, min_row=1, max_row=1))
            if not rows:
                return
            headers = [str(h or "").strip() for h in rows[0]]
        except Exception as exc:
            self._log(f"Upload columns load error: {exc}")
            return
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

        keep_checked = {
            h for h, var in self._upload_col_vars.items() if var.get()
        }

        self._upload_headers = [h for h in headers if h]
        self._upload_col_vars.clear()
        for w in self._upload_col_inner.winfo_children():
            w.destroy()

        if not self._upload_headers:
            ttk.Label(
                self._upload_col_inner,
                text="No columns found in DB file.",
                foreground="#777",
            ).pack(anchor=tk.W, padx=4, pady=4)
            return

        required_set = self._required_upload_fields()
        for h in self._upload_headers:
            is_required = h in required_set
            checked = True if is_required else (h in keep_checked if keep_checked else True)
            var = tk.BooleanVar(value=checked)
            self._upload_col_vars[h] = var
            ttk.Checkbutton(
                self._upload_col_inner,
                text=f"{h}{' (required)' if is_required else ''}",
                variable=var,
                state="disabled" if is_required else "normal",
            ).pack(anchor=tk.W, padx=6)

        self._upload_col_inner.update_idletasks()
        try:
            stat = file_path.stat()
            self._upload_cols_file_sig = (int(stat.st_mtime_ns), int(stat.st_size))
        except Exception:
            self._upload_cols_file_sig = None
        self._log(
            f"Upload columns ready: {len(self._upload_headers)} total, "
            f"{len([h for h in self._upload_headers if h in required_set])} required locked"
        )

    def _select_all_upload_cols(self) -> None:
        required_set = self._required_upload_fields()
        for h, var in self._upload_col_vars.items():
            if h in required_set:
                var.set(True)
            else:
                var.set(True)

    def _clear_upload_cols(self) -> None:
        required_set = self._required_upload_fields()
        for h, var in self._upload_col_vars.items():
            if h in required_set:
                var.set(True)
            else:
                var.set(False)

    def _selected_upload_cols(self) -> set[str]:
        return {h for h, var in self._upload_col_vars.items() if var.get()}

    def _safe_delay_value(self) -> float:
        return max(0.0, float(self._safe_delay.get() or 0.0))

    def _raise_delay_after_fail(self, reason: str, retry_after: float | None = None) -> None:
        target = 0.5
        current = self._safe_delay_value()
        changed = False
        if current < target:
            self._safe_delay.set(target)
            changed = True

        if retry_after is not None and retry_after > 0:
            self._log(
                f"Shopify requested backoff {retry_after:.2f}s ({reason}). "
                f"Waiting as requested, then continuing with delay {self._safe_delay_value():.2f}s."
            )

        if changed:
            self._log(
                f"Fail detected: {reason}. Safe Delay auto-adjusted to {target:.2f}s."
            )

    def _toggle_pause_upload(self) -> None:
        if not self._running:
            return
        self._paused = not self._paused
        if self._paused:
            self._pause_event.set()
            self._pause_btn.config(text="Resume")
            self._log("Upload paused.")
        else:
            self._pause_event.clear()
            self._pause_btn.config(text="Pause")
            self._log("Upload resumed.")

    def _stop_upload_now(self) -> None:
        if not self._running:
            return
        self._stop_event.set()
        self._pause_event.clear()
        self._paused = False
        self._pause_btn.config(text="Pause")
        self._log("Stop requested. Finishing current step...")

    def _stop_and_reset_upload(self) -> None:
        self._stop_upload_now()
        clear_session()
        self._prog["value"] = 0
        self._stats.set("Upload reset.")
        self._cont_btn.config(state="disabled")
        self._log("Upload progress/session reset requested.")

    def _payload_log_text(self, payload: dict) -> str:
        """Compact log view of outgoing Shopify payload without dumping huge base64 bodies."""
        def sanitize(value):
            if isinstance(value, dict):
                out = {}
                for key, one in value.items():
                    if key == "attachment" and isinstance(one, str):
                        out[key] = f"<base64 {len(one)} chars>"
                    else:
                        out[key] = sanitize(one)
                return out
            if isinstance(value, list):
                return [sanitize(one) for one in value]
            return value

        return json.dumps(sanitize(payload), ensure_ascii=False, default=str)

    def _scan_upload_risks(
        self,
        file_path: Path,
        selected_cols: set[str],
        start_row: int,
    ) -> tuple[dict, dict[int, str]]:
        """Pre-scan selected upload data to detect risky rows before API calls."""
        wb = None
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

        if not rows:
            return {
                "total": 0,
                "missing_id": 0,
                "duplicate_id": 0,
                "missing_image_file": 0,
                "empty_payload": 0,
                "selected_missing_headers": [],
                "readonly_selected": [],
            }, {}

        headers = [str(h or "").strip() for h in rows[0]]
        header_idx = {h: i for i, h in enumerate(headers) if h}
        data_rows = rows[1:]

        selected_missing_headers = sorted([h for h in selected_cols if h not in header_idx])
        image_src_selected = sorted(
            [h for h in selected_cols if re.match(r"^images\.\d+\.src$", h) and h in header_idx]
        )
        readonly_selected = sorted(
            [
                h
                for h in selected_cols
                if h in {"admin_graphql_api_id", "created_at", "updated_at", "published_at"}
            ]
        )

        blocked: dict[int, list[str]] = {}
        seen_ids: set[str] = set()
        missing_id = 0
        duplicate_id = 0
        missing_image_file = 0
        empty_payload = 0

        for idx in range(start_row, len(data_rows)):
            row = data_rows[idx]

            def rv_local(col: str) -> str | None:
                ci = header_idx.get(col, -1)
                val = row[ci] if 0 <= ci < len(row) else None
                return str(val).strip() if val not in (None, "") else None

            row_reasons: list[str] = []
            product_id = rv_local("id")
            if not product_id:
                missing_id += 1
                row_reasons.append("missing ID")
            else:
                if product_id in seen_ids:
                    duplicate_id += 1
                    row_reasons.append("duplicate ID")
                seen_ids.add(product_id)

            has_payload = False
            for h in selected_cols:
                if h == "id" or h not in header_idx:
                    continue
                ci = header_idx[h]
                if ci < len(row) and row[ci] not in (None, ""):
                    has_payload = True
                    break
            if not has_payload:
                empty_payload += 1

            for h in image_src_selected:
                ci = header_idx[h]
                if ci >= len(row) or row[ci] in (None, ""):
                    continue
                src = str(row[ci]).strip()
                if src.lower().startswith("http://") or src.lower().startswith("https://"):
                    continue
                local_path = Path(src)
                if not local_path.is_absolute():
                    local_path = (Path.cwd() / local_path).resolve()
                if not local_path.exists():
                    missing_image_file += 1
                    row_reasons.append(f"image file missing ({h})")

            if row_reasons:
                blocked[idx] = row_reasons

        blocked_reason_text = {idx: "; ".join(reasons) for idx, reasons in blocked.items()}
        summary = {
            "total": len(data_rows),
            "missing_id": missing_id,
            "duplicate_id": duplicate_id,
            "missing_image_file": missing_image_file,
            "empty_payload": empty_payload,
            "selected_missing_headers": selected_missing_headers,
            "readonly_selected": readonly_selected,
        }
        return summary, blocked_reason_text

    # ──────────────────────────────────────────────────
    def _start_upload(self, resume: bool = False) -> None:
        if self._running:
            messagebox.showwarning("Busy", "Upload already in progress.", parent=self)
            return

        store, token = self._app.get_conn()
        if not store:
            return

        file_path = Path(DATABASE_FILE)
        if resume:
            sess = load_session()
            if not sess:
                messagebox.showwarning("No Session", "No saved upload session found.", parent=self)
                return
            start_row = sess["pointer"]
        else:
            start_row = 0

        if not file_path.exists():
            messagebox.showwarning("File Not Found", f"Cannot find:\n{file_path}", parent=self)
            return

        selected_cols = self._selected_upload_cols()
        if not selected_cols:
            messagebox.showwarning(
                "No Upload Fields",
                "Please select at least one upload field before running upload.",
                parent=self,
            )
            return

        try:
            scan, blocked_rows = self._scan_upload_risks(file_path, selected_cols, start_row)
        except Exception as exc:
            messagebox.showerror("Pre-Scan Error", f"Cannot scan upload file:\n{exc}", parent=self)
            return

        if scan["selected_missing_headers"]:
            self._log(
                "Pre-scan warning: selected columns not found in file: "
                + ", ".join(scan["selected_missing_headers"])
            )
        if scan["readonly_selected"]:
            self._log(
                "Pre-scan warning: possible read-only fields selected: "
                + ", ".join(scan["readonly_selected"])
            )

        hazard_count = len(blocked_rows)
        if hazard_count:
            hazard_msg = (
                f"Pre-scan found {hazard_count} risky row(s).\n\n"
                f"- Missing ID: {scan['missing_id']}\n"
                f"- Duplicate ID: {scan['duplicate_id']}\n"
                f"- Missing local image file: {scan['missing_image_file']}\n"
                f"- Rows with no payload value (will skip): {scan['empty_payload']}\n\n"
                "Yes = Skip risky rows and continue upload\n"
                "No = End upload now"
            )
            go_on = messagebox.askyesno("Pre-Upload Scan", hazard_msg, parent=self)
            if not go_on:
                self._log("Upload cancelled after pre-scan.")
                return
            self._log(f"Pre-scan: continuing with {hazard_count} risky row(s) auto-skipped.")
        else:
            self._log(
                "Pre-scan: no critical risks found. "
                f"Rows with no payload value may still be skipped: {scan['empty_payload']}"
            )

        msg = f"Update Shopify products from:\n{file_path.name}"
        if resume:
            msg += f"\n\nResuming from data row {start_row + 1}."
        msg += f"\n\nSelected upload columns: {len(selected_cols)}"
        msg += f"\nPre-scan risky rows: {len(blocked_rows)}"
        msg += f"\nSafe delay: {self._safe_delay_value():.2f}s"
        msg += "\n\nContinue?"
        if not self._app.confirm_danger("Confirm Upload", msg):
            return

        self._running = True
        self._cont_btn.config(state="disabled")
        self._pause_btn.config(state="normal", text="Pause")
        self._stop_btn.config(state="normal")
        self._reset_run_btn.config(state="normal")
        self._pause_event.clear()
        self._stop_event.clear()
        self._paused = False

        def task():
            self._app.start_prog("indeterminate", use_busy_cursor=False)
            self._log(f"{'Resuming' if resume else 'Starting'} upload: {file_path.name}")
            ok_count = fail_count = skip_count = 0
            req_count = 0
            last_sku_secs = 0.0
            wb = None
            sent_count = 0
            done_count = 0
            throttle_until = 0.0
            throttle_events: list[tuple[float, str]] = []
            throttle_lock = threading.Lock()

            def set_stats(row_no: int) -> None:
                safe_delay = self._safe_delay_value()
                self._stats.set(
                    f"Row {row_no} / {total + 1}   Sent {sent_count}   Done {done_count}/{total}"
                    f"   ✓ {ok_count}   ✗ {fail_count}   - {skip_count}"
                    f"   Req {req_count}   Delay {safe_delay:.2f}s   SKU {last_sku_secs:.2f}s"
                )

            def on_backoff(wait_secs: float, reason: str) -> None:
                with throttle_lock:
                    throttle_events.append((max(0.0, float(wait_secs or 0.0)), str(reason or "throttle")))

            def flush_backoff_events() -> None:
                nonlocal throttle_until
                pending: list[tuple[float, str]] = []
                with throttle_lock:
                    if throttle_events:
                        pending = throttle_events[:]
                        throttle_events.clear()

                if not pending:
                    return

                now = time.time()
                for wait_secs, reason in pending:
                    if wait_secs > 0:
                        throttle_until = max(throttle_until, now + wait_secs)
                    self._raise_delay_after_fail(reason, wait_secs)

            try:
                wb   = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                ws   = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    self._log("ERROR: File is empty.")
                    return

                headers   = [str(h or "").strip() for h in rows[0]]
                data_rows = rows[1:]
                total     = len(data_rows)

                if "id" not in headers:
                    self._log("ERROR: 'id' column not found.")
                    messagebox.showerror(
                        "Error",
                        "The Excel file must contain an 'id' column.",
                        parent=self,
                    )
                    return

                self._log(f"Total rows: {total}. Starting from row {start_row + 1}.")
                self._log(f"Selected upload columns: {len(selected_cols)}")
                self._prog.config(maximum=total)
                self._prog["value"] = start_row
                self._app.start_prog("determinate", use_busy_cursor=False)

                max_in_flight = 3
                next_idx = start_row
                completed_rows: set[int] = set()
                resume_pointer = start_row

                def mark_done(row_idx: int) -> None:
                    nonlocal done_count, resume_pointer
                    completed_rows.add(row_idx)
                    done_count += 1
                    while resume_pointer in completed_rows:
                        resume_pointer += 1
                    save_session({"file": str(file_path), "pointer": resume_pointer, "total": total})
                    self._prog["value"] = done_count
                    set_stats(row_idx + 2)

                with ThreadPoolExecutor(max_workers=max_in_flight) as pool:
                    in_flight: dict = {}

                    while next_idx < total or in_flight:
                        flush_backoff_events()

                        if self._stop_event.is_set() and not in_flight:
                            self._log("Upload stopped by user.")
                            break

                        now = time.time()
                        throttled = now < throttle_until
                        if throttled:
                            remain = max(0.0, throttle_until - now)
                            self._stats.set(
                                f"Shopify rate-limit backoff: waiting {remain:.2f}s before next submit.  "
                                f"Done {done_count}/{total}"
                            )

                        can_submit = (
                            not self._stop_event.is_set()
                            and not self._pause_event.is_set()
                            and not throttled
                        )

                        while can_submit and next_idx < total and len(in_flight) < max_in_flight:
                            row = data_rows[next_idx]
                            idx = next_idx
                            next_idx += 1
                            row_started = time.perf_counter()

                            def rv(col: str) -> str | None:
                                i = headers.index(col) if col in headers else -1
                                val = row[i] if 0 <= i < len(row) else None
                                return str(val) if val not in (None, "") else None

                            if idx in blocked_rows:
                                skip_count += 1
                                self._log(f"Row {idx + 2}: skip (pre-scan risk: {blocked_rows[idx]})")
                                last_sku_secs = max(0.0, time.perf_counter() - row_started)
                                mark_done(idx)
                                continue

                            product_id = rv("id")
                            if not product_id:
                                skip_count += 1
                                self._log(f"Row {idx + 2}: skip (no ID)")
                                last_sku_secs = max(0.0, time.perf_counter() - row_started)
                                mark_done(idx)
                                continue

                            prod: dict = {"id": str(product_id)}
                            for ci, h in enumerate(headers):
                                if h not in selected_cols or h == "id":
                                    continue
                                if re.match(r"^images\.\d+\.(src|alt)$", h):
                                    continue
                                if ci >= len(row) or row[ci] in (None, ""):
                                    continue
                                self._set_path_value(prod, h, str(row[ci]))

                            image_map: dict[int, dict] = {}
                            for ci, h in enumerate(headers):
                                m = re.match(r"^images\.(\d+)\.(src|alt)$", h)
                                if m and h in selected_cols and ci < len(row) and row[ci] not in (None, ""):
                                    ii = int(m.group(1))
                                    ik = m.group(2)
                                    image_map.setdefault(ii, {})[ik] = str(row[ci]).strip()

                            if image_map:
                                images_payload = []
                                for _, img in sorted(image_map.items()):
                                    src = img.get("src", "").strip()
                                    if not src:
                                        continue
                                    one: dict[str, str] = {}
                                    if src.lower().startswith("http://") or src.lower().startswith("https://"):
                                        one["src"] = src
                                    else:
                                        local_path = Path(src)
                                        if not local_path.is_absolute():
                                            local_path = (Path.cwd() / local_path).resolve()
                                        if not local_path.exists():
                                            self._log(f"Row {idx + 2}: image file not found -> {local_path}")
                                            continue
                                        try:
                                            encoded = base64.b64encode(local_path.read_bytes()).decode("ascii")
                                            one["attachment"] = encoded
                                            one["filename"] = local_path.name
                                        except Exception as exc:
                                            self._log(
                                                f"Row {idx + 2}: cannot read image {local_path.name} ({exc})"
                                            )
                                            continue

                                    alt = img.get("alt", "").strip()
                                    if alt:
                                        one["alt"] = alt
                                    images_payload.append(one)

                                if images_payload:
                                    prod["images"] = images_payload

                            selected_has_image_fields = any(
                                re.match(r"^images\.\d+\.(src|alt)$", h) for h in selected_cols
                            )
                            selected_image_src_fields = sorted(
                                [h for h in selected_cols if re.match(r"^images\.\d+\.src$", h)]
                            )
                            if selected_has_image_fields and "images" not in prod:
                                sku = rv("variants.0.sku") or rv("sku") or DEFAULT_SKU
                                images_payload = self._build_images_payload_from_folder(sku, selected_image_src_fields)
                                if images_payload:
                                    prod["images"] = images_payload

                            if len(prod) == 1:
                                skip_count += 1
                                self._log(
                                    f"Row {idx + 2}  ID {product_id}: skip (no selected fields with value)"
                                )
                                last_sku_secs = max(0.0, time.perf_counter() - row_started)
                                mark_done(idx)
                                continue

                            request_payload = {"product": prod}

                            # Detect what will actually be sent to determine request count.
                            _p_keys = {k for k in prod if k not in {"id", "variants", "images", "metafields"} and prod[k] not in (None, "")}
                            _has_product = bool(_p_keys)
                            _has_variant = isinstance(prod.get("variants"), list) and any(
                                len({k: v for k, v in v.items() if k != "id" and v not in (None, "")}) > 0
                                for v in prod.get("variants", []) if isinstance(v, dict) and v.get("id")
                            )
                            _has_images = isinstance(prod.get("images"), list) and bool(prod["images"])
                            _has_meta = bool(prod.get("metafields"))
                            _mode_parts = []
                            if _has_product: _mode_parts.append("product")
                            if _has_variant: _mode_parts.append("variant")
                            if _has_images: _mode_parts.append("images")
                            if _has_meta: _mode_parts.append("metafields")
                            _mode_label = "+".join(_mode_parts) if _mode_parts else "?"
                            _req_est = len(_mode_parts)  # rough est: 1 call per resource type
                            self._log(
                                f"Row {idx + 2}  ID {product_id}: [{_mode_label}] ~{_req_est} req -> {self._payload_log_text(request_payload)}"
                            )
                            row_api_stats: dict[str, int] = {}
                            fut = pool.submit(
                                update_product_api,
                                store,
                                token,
                                str(product_id),
                                {"product": prod},
                                5,
                                row_api_stats,
                                on_backoff,
                            )
                            in_flight[fut] = (idx, str(product_id), time.perf_counter(), row_api_stats)
                            sent_count += 1
                            set_stats(idx + 2)

                            send_delay = self._safe_delay_value()
                            if send_delay > 0:
                                time.sleep(send_delay)

                        if not in_flight:
                            if self._pause_event.is_set() and not self._stop_event.is_set():
                                time.sleep(0.1)
                                continue
                            if next_idx >= total:
                                break
                            time.sleep(0.05)
                            continue

                        done, _ = wait(list(in_flight.keys()), timeout=0.2, return_when=FIRST_COMPLETED)
                        for fut in done:
                            idx, product_id, row_started, row_api_stats = in_flight.pop(fut)
                            row_reqs = int(row_api_stats.get("requests", 0))
                            req_count += row_reqs
                            last_sku_secs = max(0.0, time.perf_counter() - row_started)
                            per_req_secs = (last_sku_secs / row_reqs) if row_reqs > 0 else last_sku_secs
                            try:
                                ok, code, err = fut.result()
                            except Exception as exc:
                                ok, code, err = False, 0, str(exc)

                            if ok:
                                ok_count += 1
                                self._log(
                                    f"Row {idx + 2}  ID {product_id}: OK {code}"
                                    f"  |  req: {row_reqs}"
                                    f"  |  {per_req_secs:.2f}s/req"
                                    f"  |  {last_sku_secs:.2f}s/sku"
                                )
                            else:
                                fail_count += 1
                                retry_wait = None
                                m = re.search(r"retry[-_ ]?after\s*[:=]?\s*(\d+(?:\.\d+)?)", str(err), re.I)
                                if m:
                                    try:
                                        retry_wait = float(m.group(1))
                                    except Exception:
                                        retry_wait = None
                                self._raise_delay_after_fail(f"HTTP {code}", retry_wait)
                                self._log(
                                    f"Row {idx + 2}  ID {product_id}: FAIL {code} — {err}"
                                    f"  |  req: {row_reqs}"
                                    f"  |  {per_req_secs:.2f}s/req"
                                    f"  |  {last_sku_secs:.2f}s/sku"
                                )

                            mark_done(idx)
                            flush_backoff_events()

                # Done
                if self._stop_event.is_set():
                    self._app.set_status("Upload stopped")
                    self._cont_btn.config(state="normal")
                    self._stats.set(
                        f"Stopped — {ok_count} ok, {fail_count} failed, {skip_count} skipped"
                        f"   Req {req_count}   Delay {self._safe_delay_value():.2f}s   Last SKU {last_sku_secs:.2f}s"
                    )
                else:
                    self._log(f"Upload complete — Success: {ok_count}, Failed: {fail_count}, Skipped: {skip_count}")
                    self._app.set_status(f"Upload done — {ok_count} ok, {fail_count} failed, {skip_count} skipped")
                    self._app.refresh_live_database()
                    clear_session()
                    self._cont_btn.config(state="disabled")
                    self._stats_label.config(foreground="#333")
                    self._stats.set(
                        f"Finished — {ok_count} ok, {fail_count} failed, {skip_count} skipped"
                        f"   Req {req_count}   Delay {self._safe_delay_value():.2f}s   Last SKU {last_sku_secs:.2f}s"
                    )
                    messagebox.showinfo(
                        "Upload Complete",
                        f"Done!\n\nSuccess: {ok_count}\nFailed:  {fail_count}\nSkipped: {skip_count}",
                        parent=self,
                    )

            except Exception as exc:
                self._log(f"RUNTIME ERROR: {exc}")
                self._app.set_status("Upload interrupted — progress saved")
                self._cont_btn.config(state="normal")
                messagebox.showerror(
                    "Upload Interrupted",
                    f"Error:\n{exc}\n\n"
                    "Progress was saved.\n"
                    "Click 'Continue Upload' to resume.",
                    parent=self,
                )
            finally:
                if wb is not None:
                    try:
                        wb.close()
                    except Exception:
                        pass
                self._running = False
                self._pause_event.clear()
                self._stop_event.clear()
                self._paused = False
                self._pause_btn.config(state="disabled", text="Pause")
                self._stop_btn.config(state="disabled")
                self._reset_run_btn.config(state="disabled")
                self._app.stop_prog(clear_busy_cursor=False)

        threading.Thread(target=task, daemon=True).start()

    def _continue_upload(self) -> None:
        self._start_upload(resume=True)

    def _clear_session(self) -> None:
        if self._app.confirm_danger("Clear Session", "Clear the saved upload session?"):
            clear_session()
            self._cont_btn.config(state="disabled")
            self._stats.set("")
            self._stats_label.config(foreground="#333")
            self._log("Session cleared.")
