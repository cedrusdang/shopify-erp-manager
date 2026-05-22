"""Image replacing tool: delete existing images then upload replacements per SKU."""

from __future__ import annotations

import base64
import json
import logging
import re
import threading
import time
from pathlib import Path
from tkinter import scrolledtext, filedialog
import tkinter as tk
from tkinter import ttk, messagebox

import openpyxl

from ..api import (
    get_metafield_definition_type,
    update_product_api,
    upload_file_to_shopify_files,
    delete_product_image_by_field,
    clear_product_metafield,
)
from ..constants import DATABASE_FILE, DEFAULT_SKU, IMAGE_DIR

logger = logging.getLogger(__name__)


class ImagesReplacingTab(ttk.Frame):
    """Replace images per SKU by deleting targets then uploading local images."""

    def __init__(self, parent: ttk.Notebook, app):
        super().__init__(parent)
        self._app = app
        self._running = False
        self._headers: list[str] = []
        self._field_pick = tk.StringVar(value="")
        self._sku_mode = tk.StringVar(value="all")
        self._selected_folder = tk.StringVar(value=self._load_selected_folder())
        self._build()

    def _initial_image_folder(self) -> str:
        root = Path(IMAGE_DIR)
        if not root.is_absolute():
            root = Path.cwd() / root
        return str(root.resolve())

    def _folder_state_path(self) -> Path:
        return Path.cwd() / "image_upload_folder_state.json"

    def _load_selected_folder(self) -> str:
        default_folder = self._initial_image_folder()
        path = self._folder_state_path()
        if not path.exists():
            return default_folder
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            logger.exception("Failed to read folder state")
            return default_folder

        folder = str(data.get("selected_folder", "")).strip()
        return folder if folder else default_folder

    def _set_selected_folder(self, folder: str) -> None:
        resolved = str(Path(folder).expanduser().resolve())
        self._selected_folder.set(resolved)
        try:
            self._folder_state_path().write_text(json.dumps({"selected_folder": resolved}, ensure_ascii=True, indent=2), encoding="utf-8")
        except Exception:
            logger.exception("Failed to save folder state")

    def _build(self) -> None:
        ttk.Label(
            self,
            text="Image Replacing — delete then upload images per SKU",
            font=("Segoe UI", 10, "bold"),
            foreground="#0f4e6a",
        ).pack(anchor=tk.W, padx=10, pady=(10, 6))

        info = ttk.LabelFrame(self, text=" How It Works ", padding=8)
        info.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(
            info,
            text=(
                "1. Select a folder containing image files named by SKU (e.g., ABC123.jpg).\n"
                "2. Select a target field from the DB headers.\n"
                "3. For each SKU: delete the existing target (when supported), then upload images from the folder.\n"
                "4. Variant metafields are overwritten but not explicitly deleted prior to upload.\n"
                "5. This tab combines delete + upload to perform a replace operation."
            ),
            justify=tk.LEFT,
            foreground="#444",
            font=("Segoe UI", 9),
        ).pack(anchor=tk.W)

        options = ttk.LabelFrame(self, text=" Replace Options ", padding=8)
        options.pack(fill=tk.X, padx=8, pady=(4, 2))

        # Folder selection
        folder_row = ttk.Frame(options)
        folder_row.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(folder_row, text="Image Folder:").pack(side=tk.LEFT)
        ttk.Entry(folder_row, textvariable=self._selected_folder, state="readonly").pack(side=tk.LEFT, padx=(6, 6), fill=tk.X, expand=True)
        ttk.Button(folder_row, text="Browse...", command=self._select_folder).pack(side=tk.LEFT)

        # Field selection
        field_row = ttk.Frame(options)
        field_row.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(field_row, text="Target Field:").pack(side=tk.LEFT)
        self._field_combo = ttk.Combobox(field_row, textvariable=self._field_pick, state="readonly", width=42)
        self._field_combo.pack(side=tk.LEFT, padx=(6, 0), fill=tk.X, expand=True)
        ttk.Button(field_row, text="Refresh Fields", command=self._load_field_options).pack(side=tk.LEFT, padx=(6, 0))

        sku_row = ttk.Frame(options)
        sku_row.pack(fill=tk.X, pady=(0, 2))
        ttk.Label(sku_row, text="SKU Scope:").pack(side=tk.LEFT)
        ttk.Radiobutton(sku_row, text="All", value="all", variable=self._sku_mode).pack(side=tk.LEFT, padx=(6, 8))
        ttk.Radiobutton(sku_row, text="Single", value="single", variable=self._sku_mode).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Radiobutton(sku_row, text="Group", value="group", variable=self._sku_mode).pack(side=tk.LEFT, padx=(0, 8))

        single_sku_row = ttk.Frame(options)
        single_sku_row.pack(fill=tk.X, pady=(0, 2))
        ttk.Label(single_sku_row, text="Single SKU:").pack(side=tk.LEFT)
        self._sku_pick = tk.StringVar(value="")
        self._sku_combo = ttk.Combobox(single_sku_row, textvariable=self._sku_pick, state="disabled", width=36)
        self._sku_combo.pack(side=tk.LEFT, padx=(6, 0), fill=tk.X, expand=True)

        group_sku_row = ttk.Frame(options)
        group_sku_row.pack(fill=tk.X, pady=(0, 4))
        ttk.Label(group_sku_row, text="Group SKU (comma/newline):").pack(side=tk.LEFT)
        self._sku_group = tk.StringVar(value="")
        self._sku_group_entry = ttk.Entry(group_sku_row, textvariable=self._sku_group, state="disabled")
        self._sku_group_entry.pack(side=tk.LEFT, padx=(6, 0), fill=tk.X, expand=True)

        btns = ttk.Frame(self)
        btns.pack(fill=tk.X, padx=8, pady=4)
        ttk.Button(btns, text="Replace Images", style="Primary.TButton", command=self._start_replace).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(btns, text="Scan Folder", command=self._scan_available_images).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(btns, text="Open Folder", command=self._open_selected_folder).pack(side=tk.LEFT)

        self._stats = tk.StringVar(value="")
        ttk.Label(self, textvariable=self._stats, foreground="#333").pack(anchor=tk.W, padx=10, pady=(0, 4))

        self._prog = ttk.Progressbar(self, mode="determinate")
        self._prog.pack(fill=tk.X, padx=8, pady=(0, 6))

        ttk.Label(self, text="Available Images in Folder:", font=("Segoe UI", 9, "bold")).pack(anchor=tk.W, padx=8)
        tree_wrap = ttk.Frame(self)
        tree_wrap.pack(fill=tk.BOTH, expand=False, padx=8, pady=(2, 6))
        self._tree = ttk.Treeview(
            tree_wrap,
            columns=("filename", "size", "sku"),
            show="headings",
            height=6,
        )
        self._tree.heading("filename", text="Filename")
        self._tree.heading("size", text="Size (KB)")
        self._tree.heading("sku", text="SKU")
        self._tree.column("filename", width=280, anchor=tk.W)
        self._tree.column("size", width=80, anchor=tk.CENTER)
        self._tree.column("sku", width=120, anchor=tk.CENTER)
        ysb = ttk.Scrollbar(tree_wrap, orient=tk.VERTICAL, command=self._tree.yview)
        self._tree.configure(yscrollcommand=ysb.set)
        self._tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ysb.pack(side=tk.RIGHT, fill=tk.Y)

        ttk.Label(self, text="Replace Log:", font=("Segoe UI", 9, "bold")).pack(anchor=tk.W, padx=8)
        self._log_box = scrolledtext.ScrolledText(self, state="disabled", font=("Consolas", 8))
        self._log_box.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        self._scan_available_images()
        self._load_field_options()
        self._load_skus_from_db()

    def _log(self, msg: str) -> None:
        logger.info(f"[ImagesReplacingTab] {msg}")
        try:
            self.after(0, lambda m=msg: self._app.log(self._log_box, m))
        except Exception:
            logger.debug("[ImagesReplacingTab] UI log dispatch skipped", exc_info=True)

    def _select_folder(self) -> None:
        current = self._selected_folder.get().strip()
        initial = current if current and Path(current).is_dir() else self._initial_image_folder()
        folder = filedialog.askdirectory(title="Select Image Folder", initialdir=initial)
        if folder:
            self._set_selected_folder(folder)
            self._scan_available_images()

    def _open_selected_folder(self) -> None:
        folder = self._selected_folder.get().strip()
        if not folder:
            folder = self._initial_image_folder()
        if not folder:
            messagebox.showwarning("No Folder", "No image folder available.", parent=self)
            return
        if not Path(folder).is_dir():
            messagebox.showerror("Invalid Folder", f"Folder not found: {folder}", parent=self)
            return
        try:
            import subprocess
            import platform
            if platform.system() == "Windows":
                subprocess.Popen(["explorer", folder])
            elif platform.system() == "Darwin":
                subprocess.Popen(["open", folder])
            else:
                subprocess.Popen(["xdg-open", folder])
        except Exception as e:
            logger.exception("Failed to open folder")
            messagebox.showerror("Error", f"Could not open folder: {e}", parent=self)

    def _load_skus_from_db(self) -> None:
        wb = None
        try:
            wb = openpyxl.load_workbook(DATABASE_FILE, data_only=True)
            ws = wb.active
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            headers = [str(h or "").strip() for h in (header_row or [])]
            sku_idx = self._find_col_idx(headers, {"variants.0.sku", "sku"})
            if sku_idx is None:
                self._log("SKU column not found in database.")
                self._sku_combo["values"] = []
                return

            skus = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                sku = str(row[sku_idx]).strip() if len(row) > sku_idx and row[sku_idx] else ""
                if sku and sku != DEFAULT_SKU:
                    skus.append(sku)
            self._sku_combo["values"] = skus
        except Exception as e:
            logger.exception("Failed to load SKUs from DB")
            self._log(f"Error loading SKUs: {e}")
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

    def _find_col_idx(self, headers: list[str], accepted: set[str]) -> int | None:
        lowered = [h.lower() for h in headers]
        for name in accepted:
            key = name.lower()
            if key in lowered:
                return lowered.index(key)
        return None

    def _scan_available_images(self) -> None:
        folder = self._selected_folder.get().strip()
        if not folder:
            folder = self._initial_image_folder()

        self._tree.delete(*self._tree.get_children())

        if not folder:
            self._log("No folder available.")
            return

        folder_path = Path(folder)
        if not folder_path.is_dir():
            self._log(f"Folder not found: {folder}")
            return

        self._log(f"Scanning folder: {folder}")

        try:
            images = []
            for file_path in folder_path.glob("*"):
                if file_path.is_file() and file_path.suffix.lower() in {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}:
                    stem = file_path.stem
                    sku = (stem or "").strip()
                    if not sku:
                        continue
                    size_kb = file_path.stat().st_size / 1024
                    images.append((file_path.name, size_kb, sku))
                    self._tree.insert("", "end", values=(file_path.name, f"{size_kb:.1f}", sku))

            self._log(f"Found {len(images)} image(s) in folder.")
            self._stats.set(f"Available: {len(images)} image(s)")
        except Exception as e:
            logger.exception("Folder scan failed")
            self._log(f"Error scanning folder: {e}")

    def _load_field_options(self) -> None:
        try:
            wb = openpyxl.load_workbook(DATABASE_FILE, read_only=True, data_only=True)
            ws = wb.active
            rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
            first = next(rows, None)
            headers = [str(h or "").strip() for h in (first or [])]
            wb.close()
        except Exception as exc:
            logger.exception("Failed to load fields from DB")
            self._field_combo["values"] = []
            self._field_pick.set("")
            self._field_combo.config(state="disabled")
            self._log(f"Error loading fields: {exc}")
            return

        self._headers = headers
        fields = [h for h in headers if h and (h.startswith("metafields.") or h == "images" or (h.startswith("images.") and h.endswith(".src")) or h.startswith("variants.0.metafields."))]
        self._field_combo["values"] = fields
        current = self._field_pick.get().strip()
        if current in fields:
            self._field_pick.set(current)
            self._field_combo.config(state="readonly")
        elif fields:
            self._field_pick.set(fields[0])
            self._field_combo.config(state="readonly")
        else:
            self._field_pick.set("")
            self._field_combo.config(state="disabled")
            self._log("No supported image target fields found in the database header.")

    def _selected_sku_set(self) -> set[str]:
        mode = self._sku_mode.get()
        if mode == "single":
            sku = self._sku_pick.get().strip()
            return {sku} if sku else set()
        elif mode == "group":
            group_text = self._sku_group.get().strip()
            if not group_text:
                return set()
            skus = re.split(r"[,\n]+", group_text)
            return {s.strip() for s in skus if s.strip()}
        else:
            return set()

    def _is_product_image_field(self, field_name: str) -> bool:
        return field_name == "images" or (field_name.startswith("images.") and field_name.endswith(".src"))

    def _is_product_metafield_field(self, field_name: str) -> bool:
        return field_name.startswith("metafields.")

    def _is_variant_metafield_field(self, field_name: str) -> bool:
        return field_name.startswith("variants.0.metafields.")

    def _upload_product_image(self, store: str, token: str, product_id: str, sku: str, img_path: Path) -> tuple[bool, str, str]:
        with open(img_path, "rb") as f:
            img_data = f.read()
        attachment = base64.b64encode(img_data).decode("utf-8")

        payload = {
            "image": {
                "attachment": attachment,
                "alt": f"Product image: {sku}",
            }
        }

        url = f"https://{store}/admin/api/2024-10/products/{product_id}/images.json"
        import requests
        resp = requests.post(
            url,
            headers={"X-Shopify-Access-Token": token, "Content-Type": "application/json"},
            json=payload,
            timeout=45,
        )
        if not resp.ok:
            error_msg = resp.text[:300] if resp.text else f"HTTP {resp.status_code}"
            return False, "", error_msg

        try:
            body = resp.json()
        except ValueError:
            body = {}
        image = body.get("image", {}) if isinstance(body, dict) else {}
        uploaded_src = str(image.get("src", "")).strip()
        return True, uploaded_src, ""

    def _start_replace(self) -> None:
        folder = self._selected_folder.get().strip()
        field_name = self._field_pick.get().strip()

        if not folder:
            messagebox.showwarning("No Folder", "Please select an image folder.", parent=self)
            return

        if not field_name:
            messagebox.showwarning("No Field", "Please select target field first.", parent=self)
            return

        if not Path(folder).is_dir():
            messagebox.showerror("Invalid Folder", f"Folder not found: {folder}", parent=self)
            return

        store, token = self._app.get_conn()
        if not store or not token:
            return

        sku_set = self._selected_sku_set()
        scope_desc = f"({len(sku_set)} SKUs)" if sku_set else "(all SKUs)"
        msg = (
            f"Replace images from:\n{folder}\n\n"
            f"Field: {field_name}\n"
            f"Scope: {scope_desc}\n\n"
            "This will attempt to delete the existing target (when supported) and then upload replacements. Continue?"
        )
        if not self._app.confirm("Confirm Image Replace", msg):
            return

        self._running = True
        self._app.start_prog("determinate", use_busy_cursor=False)
        self._log("Replace task started…")

        def task():
            try:
                self._run_replace_task(store, token, folder, sku_set)
            except Exception as exc:
                logger.exception("Replace task crashed")
                self._ui_async(lambda: self._app.stop_prog(clear_busy_cursor=False))
                self._ui_async(lambda: messagebox.showerror("Replace Error", str(exc), parent=self))
            finally:
                self._running = False

        thread = threading.Thread(target=task, daemon=True)
        thread.start()

    def _ui_async(self, fn) -> None:
        try:
            self.after(0, fn)
        except Exception:
            logger.debug("[ImagesReplacingTab] UI async dispatch failed", exc_info=True)

    def _run_replace_task(self, store: str, token: str, folder: str, sku_set: set[str]) -> None:
        start_time = time.time()
        wb = None
        field_name = self._field_pick.get().strip()

        try:
            # Load DB SKU→ID map
            self._log("Loading product database")
            wb = openpyxl.load_workbook(DATABASE_FILE, data_only=True)
            ws = wb.active
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            headers = [str(h or "").strip() for h in (header_row or [])]
            id_idx = self._find_col_idx(headers, {"id"})
            sku_idx = self._find_col_idx(headers, {"variants.0.sku", "sku"})
            variant_id_idx = self._find_col_idx(headers, {"variants.0.id"})
            if id_idx is None or sku_idx is None:
                raise RuntimeError("Database must contain 'id' and 'variants.0.sku' (or 'sku') columns.")

            sku_to_id = {}
            sku_to_variant_id: dict[str, str] = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(row) <= max(id_idx, sku_idx):
                    continue
                prod_id = str(row[id_idx]).strip() if row[id_idx] else ""
                sku = str(row[sku_idx]).strip() if row[sku_idx] else ""
                if prod_id and sku and sku != DEFAULT_SKU:
                    sku_to_id[sku] = prod_id
                    if variant_id_idx is not None and len(row) > variant_id_idx and row[variant_id_idx]:
                        sku_to_variant_id[sku] = str(row[variant_id_idx]).strip()
            self._log(f"Loaded {len(sku_to_id)} SKU→ID mappings from database")

            # Scan folder
            self._log("Scanning image folder")
            folder_path = Path(folder)
            images_by_sku: dict[str, list[Path]] = {}
            for file_path in folder_path.glob("*"):
                if file_path.is_file() and file_path.suffix.lower() in {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}:
                    stem = file_path.stem
                    sku = (stem or "").strip()
                    if not sku:
                        continue
                    images_by_sku.setdefault(sku, []).append(file_path)

            self._log(f"Found images for {len(images_by_sku)} SKU(s)")

            # Determine SKUs to process
            process_skus = sku_set if sku_set else set(images_by_sku.keys())
            process_skus = process_skus & set(sku_to_id.keys())

            if not process_skus:
                self._ui_async(lambda: self._log("No matching SKUs found in database."))
                self._ui_async(lambda: self._app.stop_prog(clear_busy_cursor=False))
                return

            self._log(f"Will process {len(process_skus)} SKU(s)")
            ok_count = 0
            fail_count = 0

            for i, sku in enumerate(sorted(process_skus)):
                if not self._running:
                    self._log("Replace cancelled by user")
                    break

                prod_id = sku_to_id[sku]
                sku_images = images_by_sku.get(sku, [])

                # 1) delete existing target when supported
                deleted_ok = True
                if self._is_product_image_field(field_name):
                    ok, _status, msg = delete_product_image_by_field(store, token, prod_id, field_name)
                    if not ok:
                        deleted_ok = False
                        self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: delete failed ✗ ({msg})")
                elif self._is_product_metafield_field(field_name):
                    ok, _status, msg = clear_product_metafield(store, token, prod_id, field_name)
                    if not ok:
                        deleted_ok = False
                        self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: clear metafield failed ✗ ({msg})")
                else:
                    # Variant metafields: deletion not supported here; will overwrite during upload
                    self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: skip delete (variant metafield or unsupported)")

                if not deleted_ok:
                    fail_count += 1
                    progress = int((i + 1) / len(process_skus) * 100)
                    self._ui_async(lambda p=progress: self._app.set_prog_value(p, 100))
                    continue

                # 2) upload replacements (if images exist)
                if not sku_images:
                    self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: no images to upload, skip")
                    continue

                uploaded_any = False
                for img_path in sku_images:
                    try:
                        if not img_path.is_file() or not img_path.stat().st_size:
                            self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: {img_path.name} ✗ (not readable or empty)")
                            continue

                        if self._is_product_image_field(field_name):
                            ok_upload, uploaded_value, upload_error = self._upload_product_image(store, token, prod_id, sku, img_path)
                        else:
                            ok_upload, uploaded_file_id, uploaded_file_url, upload_error = upload_file_to_shopify_files(
                                store,
                                token,
                                img_path,
                                alt=f"Product image: {sku}",
                            )
                            if ok_upload:
                                ok_value, uploaded_value, value_error = self._resolve_metafield_upload_value(store, token, field_name, uploaded_file_id, uploaded_file_url)
                                if not ok_value:
                                    self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: {img_path.name} ✗ ({value_error})")
                                    continue

                        if not ok_upload:
                            self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: {img_path.name} ✗ ({upload_error})")
                            continue

                        # If metafield target, write field
                        if not self._is_product_image_field(field_name):
                            field_payload = self._build_field_update_payload(field_name, uploaded_value, sku_to_variant_id.get(sku, ""))
                            if field_payload is None:
                                self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: {img_path.name} ✗ (unsupported field mapping: {field_name})")
                                continue
                            ok_field, _status_field, msg_field = update_product_api(store, token, prod_id, field_payload)
                            if not ok_field:
                                self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: {img_path.name} ✗ (field update failed: {msg_field})")
                                continue

                        uploaded_any = True
                        self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: {img_path.name} ✓")
                    except Exception as e:
                        self._log(f"  [{i + 1}/{len(process_skus)}] SKU {sku}: {img_path.name} ✗ ({e})")

                if uploaded_any:
                    ok_count += 1
                else:
                    fail_count += 1

                progress = int((i + 1) / len(process_skus) * 100)
                self._ui_async(lambda p=progress: self._app.set_prog_value(p, 100))

            elapsed = time.time() - start_time
            self._log(f"Replace completed in {elapsed:.1f}s (ok={ok_count}, fail={fail_count})")
            self._ui_async(lambda: self._stats.set(f"Processed: {ok_count + fail_count} | Success: {ok_count} | Failed: {fail_count}"))
            self._ui_async(lambda: self._app.stop_prog(clear_busy_cursor=False))
            self._ui_async(lambda: messagebox.showinfo("Replace Complete", f"Success: {ok_count}\nFailed: {fail_count}\nTime: {elapsed:.1f}s", parent=self))

        except Exception:
            logger.exception("Replace task failed")
            self._ui_async(lambda: self._app.stop_prog(clear_busy_cursor=False))
            raise
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

    def _resolve_metafield_upload_value(self, store: str, token: str, field_name: str, file_id: str, file_url: str) -> tuple[bool, str, str]:
        # Reuse logic similar to upload tab
        if self._is_product_metafield_field(field_name):
            parts = field_name.split(".", 2)
            if len(parts) != 3:
                return False, "", f"Invalid metafield name: {field_name}"
            _prefix, namespace, key = parts
            owner_type = "PRODUCT"
        elif self._is_variant_metafield_field(field_name):
            parts = field_name.split(".", 4)
            if len(parts) != 5:
                return False, "", f"Invalid variant metafield name: {field_name}"
            _variants, _index, _metafields, namespace, key = parts
            owner_type = "PRODUCTVARIANT"
        else:
            return True, file_url, ""

        definition_type = get_metafield_definition_type(store, token, owner_type, namespace, key)
        if definition_type == "file_reference":
            if not file_id:
                return False, "", "Shopify file upload did not return a file reference id"
            return True, file_id, ""
        if definition_type.startswith("list."):
            return False, "", f"Metafield type '{definition_type}' is not supported by this tab"
        return True, file_url, ""

    def _build_field_update_payload(self, field_name: str, value: str, variant_id: str = "") -> dict | None:
        if self._is_product_metafield_field(field_name):
            parts = field_name.split(".", 2)
            if len(parts) != 3:
                return None
            _prefix, namespace, key = parts
            return {"product": {"metafields": {namespace: {key: value}}}}

        if self._is_variant_metafield_field(field_name):
            parts = field_name.split(".", 4)
            if len(parts) != 5 or not variant_id:
                return None
            _variants, _index, _metafields, namespace, key = parts
            return {"product": {"variants": [{"id": variant_id, "metafields": {namespace: {key: value}}}]}}

        if self._is_product_image_field(field_name):
            return None

        return None
