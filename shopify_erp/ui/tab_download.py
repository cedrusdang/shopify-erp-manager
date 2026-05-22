"""
Download tab — fetch products from Shopify and save to Excel.
"""

from __future__ import annotations

import re
import json
import logging
import threading
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from collections.abc import Iterable
from tkinter import scrolledtext
import tkinter as tk
from tkinter import ttk, messagebox

import openpyxl
import requests
from openpyxl.styles import Font

from ..api import fetch_all_products, fetch_metafield_definitions, fetch_products_page, fetch_products_range, get_by_path, enrich_products_with_metafields
from ..backup import create_backup, open_file
from ..constants import DATABASE_FILE, DEFAULT_SKU, IMAGE_DIR, REQUIRED_FIELDS
from ..download_field_state import load_field_state, save_field_state
from ..field_presets import delete_preset, load_presets, upsert_preset
from ..ui.widgets import FieldSelector

logger = logging.getLogger(__name__)


def _needs_metafield_enrichment(fields: list[str]) -> bool:
    return any(
        isinstance(field, str)
        and (field.startswith("metafields.") or ".metafields." in field)
        for field in fields
    )


class DownloadTab(ttk.Frame):
    def __init__(self, parent: ttk.Notebook, app):
        super().__init__(parent)
        self._app = app
        self._running = False
        self._discover_generation = 0
        self._presets: list[dict] = []
        self._progress_nodes: dict[str, str] = {}
        self._build()

    # ──────────────────────────────────────────────────
    def _build(self) -> None:
        saved_state = load_field_state()
        saved_optional = [str(f) for f in saved_state.get("optional_fields", []) if str(f).strip()]
        saved_selected = [str(f) for f in saved_state.get("selected_fields", []) if str(f).strip()]

        # Left: field selector
        left = ttk.Frame(self)
        left.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=8, pady=8)
        ttk.Label(
            left,
            text="(3) Select fields, then click Download",
            font=("Segoe UI", 10, "bold"),
            foreground="#0f6b45",
        ).pack(anchor=tk.W, pady=(0, 6))
        # Query-first mode: optional fields are discovered dynamically.
        self.fields = FieldSelector(left, REQUIRED_FIELDS, saved_optional, on_change=self._save_field_state)
        self.fields.pack(fill=tk.BOTH, expand=True)
        if saved_selected:
            self.fields.set_selected(saved_selected)

        preset = ttk.LabelFrame(left, text=" Selection Presets ", padding=8)
        preset.pack(fill=tk.X, pady=(6, 4))
        row1 = ttk.Frame(preset)
        row1.pack(fill=tk.X)
        ttk.Label(row1, text="Preset Name:").pack(side=tk.LEFT)
        self._preset_name = tk.StringVar(value="")
        ttk.Entry(row1, textvariable=self._preset_name, width=28).pack(
            side=tk.LEFT, padx=(6, 6), fill=tk.X, expand=True
        )
        ttk.Button(row1, text="Save/Update", command=self._save_selection_preset).pack(side=tk.LEFT)

        row2 = ttk.Frame(preset)
        row2.pack(fill=tk.X, pady=(6, 0))
        ttk.Label(row2, text="Saved Presets:").pack(side=tk.LEFT)
        self._preset_pick = tk.StringVar(value="")
        self._preset_combo = ttk.Combobox(
            row2,
            textvariable=self._preset_pick,
            state="readonly",
            width=30,
        )
        self._preset_combo.pack(side=tk.LEFT, padx=(6, 6), fill=tk.X, expand=True)
        self._preset_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_pick_preset())

        row3 = ttk.Frame(preset)
        row3.pack(fill=tk.X, pady=(6, 0))
        ttk.Button(row3, text="Load", command=self._load_selection_preset).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(row3, text="Delete", command=self._delete_selection_preset).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(row3, text="Refresh", command=self._refresh_presets).pack(side=tk.LEFT)

        ttk.Button(
            left,
            text="(3a) Discover More Fields",
            command=self._discover_fields,
        ).pack(fill=tk.X, pady=4)
        ttk.Button(
            left,
            text="Reset Extra Fields",
            command=self._reset_discovered_fields,
        ).pack(fill=tk.X, pady=(0, 4))

        self._refresh_presets()

        # Right: controls
        right = ttk.LabelFrame(self, text=" Actions ", padding=10, width=240)
        right.pack(side=tk.RIGHT, fill=tk.Y, padx=(0, 8), pady=8)
        right.pack_propagate(False)

        ttk.Label(right, text="Database File (fixed):").pack(anchor=tk.W)
        ttk.Label(
            right,
            text=DATABASE_FILE,
            foreground="#0f6b45",
            font=("Segoe UI", 9, "bold"),
        ).pack(anchor=tk.W, pady=(2, 0))

        ttk.Separator(right).pack(fill=tk.X, pady=10)
        ttk.Label(right, text="Download Scope:", font=("Segoe UI", 8, "bold")).pack(anchor=tk.W)
        self._scope_mode = tk.StringVar(value="All Pages")
        self._scope_combo = ttk.Combobox(
            right,
            textvariable=self._scope_mode,
            state="readonly",
            values=["All Pages", "Single Page", "From-To"],
            width=22,
        )
        self._scope_combo.pack(fill=tk.X, pady=(2, 4))
        self._scope_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_scope_changed())

        page_wrap = ttk.Frame(right)
        page_wrap.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(page_wrap, text="Page #:").pack(side=tk.LEFT)
        self._page_no = tk.IntVar(value=1)
        self._page_spin = ttk.Spinbox(page_wrap, from_=1, to=999999, textvariable=self._page_no, width=8)
        self._page_spin.pack(side=tk.LEFT, padx=(6, 0))

        range_wrap = ttk.Frame(right)
        range_wrap.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(range_wrap, text="From:").pack(side=tk.LEFT)
        self._page_from = tk.IntVar(value=1)
        self._from_spin = ttk.Spinbox(range_wrap, from_=1, to=999999, textvariable=self._page_from, width=7)
        self._from_spin.pack(side=tk.LEFT, padx=(6, 4))
        ttk.Label(range_wrap, text="To:").pack(side=tk.LEFT)
        self._page_to = tk.IntVar(value=3)
        self._to_spin = ttk.Spinbox(range_wrap, from_=1, to=999999, textvariable=self._page_to, width=7)
        self._to_spin.pack(side=tk.LEFT, padx=(6, 0))

        self._scope_info = tk.StringVar(value="Mode: All Pages (download everything)")
        ttk.Label(right, textvariable=self._scope_info, foreground="#666", wraplength=210, font=("Segoe UI", 8)).pack(anchor=tk.W, pady=(0, 6))

        ttk.Button(
            right, text="(3b) Download Products",
            command=self._start_download,
            style="Primary.TButton",
        ).pack(fill=tk.X, pady=3)

        ttk.Label(
            right,
            text="Image download moved to 'Image Download' tab.",
            foreground="#666",
            wraplength=210,
            font=("Segoe UI", 8),
        ).pack(anchor=tk.W, pady=(6, 4))

        ttk.Button(
            right, text="📂  Open Current File",
            command=self._open_current,
        ).pack(fill=tk.X, pady=3)

        ttk.Separator(right).pack(fill=tk.X, pady=10)
        ttk.Label(
            right, text="Last Downloaded:",
            font=("Segoe UI", 8, "bold"),
        ).pack(anchor=tk.W)
        self._last_dl = tk.StringVar(value="—")
        ttk.Label(
            right, textvariable=self._last_dl,
            foreground="#555", wraplength=210, font=("Segoe UI", 8),
        ).pack(anchor=tk.W, pady=2)

        ttk.Label(
            right, text="Progress Tree:", font=("Segoe UI", 8, "bold"),
        ).pack(anchor=tk.W, pady=(8, 0))
        self._progress_tree = ttk.Treeview(
            right,
            columns=("status", "detail"),
            show="tree headings",
            height=8,
        )
        self._progress_tree.heading("#0", text="Step")
        self._progress_tree.heading("status", text="Status")
        self._progress_tree.heading("detail", text="Detail")
        self._progress_tree.column("#0", width=120, anchor=tk.W)
        self._progress_tree.column("status", width=70, anchor=tk.CENTER)
        self._progress_tree.column("detail", width=180, anchor=tk.W)
        self._progress_tree.pack(fill=tk.X, pady=(3, 6))

        ttk.Label(
            right, text="Log:", font=("Segoe UI", 8, "bold"),
        ).pack(anchor=tk.W, pady=(10, 0))
        self._log_box = scrolledtext.ScrolledText(
            right, state="disabled", font=("Consolas", 8), height=16
        )
        self._log_box.pack(fill=tk.BOTH, expand=True)
        self._refresh_image_url_field_candidates()
        self._on_scope_changed()

    def _save_field_state(self) -> None:
        save_field_state(self.fields.get_optional_fields(), self.fields.get_selected())

    def _with_auto_required_fields(self, selected_fields: list[str]) -> tuple[list[str], list[str]]:
        """Auto-append hidden dependency fields needed for reliable upload later.

        Rules:
        - Always keep `id`.
        - If any selected field targets variant payload (variants.*), include `variants.0.id`.
        """
        fields = [f for f in selected_fields if str(f).strip()]
        out = list(fields)
        added: list[str] = []

        if "id" not in out:
            out.append("id")
            added.append("id")

        needs_variant_id = any(f.startswith("variants.") and f != "variants.0.id" for f in out)
        if needs_variant_id and "variants.0.id" not in out:
            out.append("variants.0.id")
            added.append("variants.0.id")

        return out, added

    # ──────────────────────────────────────────────────
    def current_xlsx(self) -> Path:
        return Path(DATABASE_FILE)

    def _log(self, msg: str) -> None:
        self._app.log(self._log_box, msg)

    def _safe_sku(self, value: str) -> str:
        sku = re.sub(r"[^A-Za-z0-9._-]+", "_", (value or "").strip())
        return (sku[:80] or "product")

    def _field_suffix(self, field_name: str) -> str:
        m = re.match(r"^images\.(\d+)\.src$", field_name)
        if m:
            return f"img_{m.group(1)}"
        clean = re.sub(r"[^A-Za-z0-9._-]+", "_", (field_name or "").strip())
        return clean[:40] or "img"

    def _field_folder_name(self, field_name: str) -> str:
        clean = re.sub(r"[^A-Za-z0-9._-]+", "_", (field_name or "").strip())
        return clean[:80] or "images"

    def _reset_progress_tree(self, mode: str, target_desc: str) -> None:
        def _ui() -> None:
            self._progress_tree.delete(*self._progress_tree.get_children())
            run = self._progress_tree.insert("", tk.END, text="Download", values=("RUNNING", f"Mode {mode}: {target_desc}"))
            fetch = self._progress_tree.insert(run, tk.END, text="Fetch Products", values=("PENDING", "Waiting"))
            images = self._progress_tree.insert(run, tk.END, text="Images (Image Tab)", values=("PENDING", "Waiting"))
            write = self._progress_tree.insert(run, tk.END, text="Write Excel", values=("PENDING", "Waiting"))
            refresh = self._progress_tree.insert(run, tk.END, text="Refresh Live DB", values=("PENDING", "Waiting"))
            backup = self._progress_tree.insert(run, tk.END, text="Create Backup", values=("PENDING", "Waiting"))
            finish = self._progress_tree.insert(run, tk.END, text="Finish", values=("PENDING", "Waiting"))
            self._progress_nodes = {
                "run": run,
                "fetch": fetch,
                "images": images,
                "write": write,
                "refresh": refresh,
                "backup": backup,
                "finish": finish,
            }
            self._progress_tree.item(run, open=True)

        self.after(0, _ui)

    def _set_progress_step(self, key: str, status: str, detail: str = "") -> None:
        def _ui() -> None:
            iid = self._progress_nodes.get(key)
            if not iid:
                return
            self._progress_tree.set(iid, "status", status)
            self._progress_tree.set(iid, "detail", detail)
            self._progress_tree.selection_set(iid)
            self._progress_tree.see(iid)

        self.after(0, _ui)

    def _refresh_image_url_field_candidates(self) -> None:
        # Image flow was moved out of Download tab; keep as no-op for compatibility.
        return

    def _set_by_path(self, obj: object, path: str, value: object) -> bool:
        """Best-effort setter for dot path like variants.0.sku."""
        parts = path.split(".")
        if not parts:
            return False
        cur = obj
        try:
            for p in parts[:-1]:
                if p.isdigit():
                    cur = cur[int(p)]  # type: ignore[index]
                else:
                    cur = cur[p]  # type: ignore[index]
            last = parts[-1]
            if last.isdigit():
                cur[int(last)] = value  # type: ignore[index]
            else:
                cur[last] = value  # type: ignore[index]
            return True
        except Exception:
            return False

    def _extract_url_candidates(self, raw: object) -> list[str]:
        urls: list[str] = []

        def _add(x: object) -> None:
            if isinstance(x, str):
                s = x.strip()
                if s.startswith(("http://", "https://")):
                    urls.append(s)
                return
            if isinstance(x, dict):
                for key in ("src", "url"):
                    val = x.get(key)
                    if isinstance(val, str) and val.strip().startswith(("http://", "https://")):
                        urls.append(val.strip())
                return
            if isinstance(x, Iterable) and not isinstance(x, (str, bytes, bytearray)):
                for item in x:
                    _add(item)

        _add(raw)
        return urls

    def _download_images_to_local(self, products: list[dict], image_field: str) -> tuple[int, int]:
        """Download image URLs from chosen field and replace that field with local file path."""
        image_dir = Path(IMAGE_DIR)
        image_dir.mkdir(exist_ok=True)

        ok = 0
        fail = 0
        valid_ext = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
        field_suffix = self._field_suffix(image_field)
        field_folder = image_dir / self._field_folder_name(image_field)
        field_folder.mkdir(parents=True, exist_ok=True)

        for p in products:
            raw = get_by_path(p, image_field)
            src_list = self._extract_url_candidates(raw)
            if not src_list:
                continue

            sku = str(get_by_path(p, "variants.0.sku") or "").strip()
            if not sku:
                sku = DEFAULT_SKU
            safe_sku = self._safe_sku(sku)

            for i, src in enumerate(src_list):

                try:
                    ext = Path(urlparse(src).path).suffix.lower()
                    if ext not in valid_ext:
                        ext = ".jpg"
                    file_name = f"{safe_sku}_{field_suffix}{ext}"
                    target = field_folder / file_name

                    r = requests.get(src, timeout=30)
                    r.raise_for_status()
                    target.write_bytes(r.content)

                    # Replace chosen image field with the first downloaded local path.
                    if i == 0:
                        self._set_by_path(p, image_field, str(target.resolve()))
                    ok += 1
                except Exception:
                    fail += 1
                    self._log(
                        "Image download FAIL context -> "
                        + json.dumps(
                            {
                                "product_id": str(get_by_path(p, "id") or ""),
                                "sku": sku,
                                "field": image_field,
                                "url": src,
                            },
                            ensure_ascii=False,
                        )
                    )

        return ok, fail

    # ──────────────────────────────────────────────────
    def _discover_fields(self) -> None:
        store, token = self._app.get_conn()
        if not store:
            return

        self._discover_generation += 1
        discover_generation = self._discover_generation

        def task():
            self._app.start_prog(use_busy_cursor=True)
            self._log("Discovering metafield definitions…")
            try:
                mf = fetch_metafield_definitions(store, token)
                self.after(0, lambda: self._apply_discovered_fields(mf, discover_generation))
            except Exception as exc:
                self.after(0, lambda: self._on_discover_error(exc))
            finally:
                self._app.stop_prog(clear_busy_cursor=True)

        threading.Thread(target=task, daemon=True).start()

    def _apply_discovered_fields(self, mf: list[str], discover_generation: int) -> None:
        if discover_generation != self._discover_generation:
            self._log("Skipped stale discover result (field list changed meanwhile).")
            return
        added = self.fields.add_optional_fields(mf)
        self._save_field_state()
        self._refresh_image_url_field_candidates()
        mf_fields = [f for f in mf if ".metafields." in f]
        self._log(f"Found {len(mf)} definitions. {added} new fields added.")
        if mf_fields:
            self._log(f"  Metafields discovered: {', '.join(mf_fields)}")
        else:
            self._log("  WARNING: No metafield paths found in discover result.")
        self._app.set_status(f"Discovered {len(mf)} metafield definitions")

    def _on_discover_error(self, exc: Exception) -> None:
        self._log(f"Error: {exc}")
        messagebox.showerror("Error", str(exc), parent=self)

    def _reset_discovered_fields(self) -> None:
        # Invalidate any in-flight discover task so stale result cannot re-add fields.
        self._discover_generation += 1
        removed = self.fields.reset_optional_fields()
        self._save_field_state()
        self._refresh_image_url_field_candidates()
        self._log(f"Reset fields: removed {removed} discovered extra field(s).")
        self._app.set_status("Field list reset to default optional fields")

    def _on_scope_changed(self) -> None:
        single = self._scope_mode.get() == "Single Page"
        ranged = self._scope_mode.get() == "From-To"
        self._page_spin.configure(state="normal" if single else "disabled")
        self._from_spin.configure(state="normal" if ranged else "disabled")
        self._to_spin.configure(state="normal" if ranged else "disabled")
        if single:
            self._scope_info.set(f"Mode: Single Page (current page #{max(1, int(self._page_no.get() or 1))})")
        elif ranged:
            f = max(1, int(self._page_from.get() or 1))
            t = max(f, int(self._page_to.get() or f))
            self._scope_info.set(f"Mode: From-To (pages {f} to {t})")
        else:
            self._scope_info.set("Mode: All Pages (download everything)")

    def _refresh_presets(self) -> None:
        self._presets = load_presets()
        names = [p.get("name", "") for p in self._presets if p.get("name")]
        self._preset_combo["values"] = names
        current = self._preset_pick.get().strip()
        if current and current in names:
            self._preset_combo.set(current)
        elif names:
            self._preset_combo.set(names[0])
            self._preset_pick.set(names[0])
        else:
            self._preset_combo.set("")
            self._preset_pick.set("")

    def _on_pick_preset(self) -> None:
        name = self._preset_pick.get().strip()
        if name:
            self._preset_name.set(name)

    def _selected_preset_row(self) -> dict | None:
        name = self._preset_pick.get().strip()
        if not name:
            return None
        for p in self._presets:
            if p.get("name", "").strip() == name:
                return p
        return None

    def _save_selection_preset(self) -> None:
        name = self._preset_name.get().strip()
        if not name:
            messagebox.showwarning("Preset Name", "Please enter a preset name.", parent=self)
            return
        fields = self.fields.get_selected()
        upsert_preset(name, fields)
        self._refresh_presets()
        self._refresh_image_url_field_candidates()
        self._preset_pick.set(name)
        self._preset_combo.set(name)
        self._log(f"Preset saved: {name} ({len(fields)} fields)")

    def _load_selection_preset(self) -> None:
        row = self._selected_preset_row()
        if not row:
            messagebox.showwarning("No Preset", "Please select a saved preset.", parent=self)
            return
        fields = [str(f) for f in row.get("fields", [])]
        # If preset contains fields discovered previously, re-add them before applying.
        self.fields.add_optional_fields(fields)
        self.fields.set_selected(fields)
        self._save_field_state()
        self._refresh_image_url_field_candidates()
        self._log(f"Preset loaded: {row.get('name', '')} ({len(fields)} fields)")

    def _delete_selection_preset(self) -> None:
        row = self._selected_preset_row()
        if not row:
            messagebox.showwarning("No Preset", "Please select a saved preset.", parent=self)
            return
        name = str(row.get("name", ""))
        if not self._app.confirm("Delete Preset", f"Delete preset '{name}'?"):
            return
        delete_preset(name)
        self._refresh_presets()
        self._log(f"Preset deleted: {name}")

    # ──────────────────────────────────────────────────
    def _start_download(self) -> None:
        if self._running:
            messagebox.showwarning("Busy", "Download already in progress.", parent=self)
            return

        store, token = self._app.get_conn()
        if not store:
            return

        fields = self.fields.get_selected()
        self._refresh_image_url_field_candidates()
        if not fields:
            messagebox.showwarning("No Fields", "Select at least one field.", parent=self)
            return
        fields, auto_added_fields = self._with_auto_required_fields(fields)
        if auto_added_fields:
            self._log(
                "Auto-added dependency field(s): " + ", ".join(auto_added_fields)
                + " (for upload compatibility)"
            )

        out = self.current_xlsx()
        mode = self._scope_mode.get()
        single = mode == "Single Page"
        ranged = mode == "From-To"
        page_no = max(1, int(self._page_no.get() or 1))
        page_from = max(1, int(self._page_from.get() or 1))
        page_to = max(1, int(self._page_to.get() or 1))
        if ranged and page_from > page_to:
            messagebox.showwarning("Invalid Range", "From page must be <= To page.", parent=self)
            return

        if single:
            target_desc = f"page {page_no}"
        elif ranged:
            target_desc = f"pages {page_from} to {page_to}"
        else:
            target_desc = "ALL pages"
        if not self._app.confirm_danger(
            "Confirm Download",
            f"Fetch {target_desc} from Shopify and save to:\n"
            f"{out.resolve()}\n\n"
            f"Existing file will be overwritten.\n\nContinue?",
        ):
            return

        self._running = True

        def task():
            current_step = "fetch"
            self._app.start_prog("indeterminate", use_busy_cursor=False)
            self._reset_progress_tree(mode, target_desc)
            self._log("Starting download…")
            logger.info(f"Download started - Mode: {mode}, Store: {store}")
            try:
                self._set_progress_step("fetch", "RUNNING", "Fetching products")
                if single:
                    logger.info(f"Fetching single page {page_no}")
                    products, actual_page, has_next = fetch_products_page(
                        store,
                        token,
                        page_no,
                        on_progress=lambda m: (self._log(m), self._app.set_status(m)),
                    )
                    total = len(products)
                    if actual_page != page_no and not products:
                        raise ValueError(
                            f"Requested page {page_no} not found. Last available page is {actual_page}."
                        )
                    if page_no == 1 and total == 0:
                        self._log(
                            "WARNING: Page 1 returned 0 products. Check app scope 'read_products' and verify this store has products matching current access."
                        )
                    self._log(
                        f"Fetched page {actual_page} ({total} products). Next page exists: {'Yes' if has_next else 'No'}"
                    )
                    logger.info(f"Single page download success: page {actual_page}, {total} products")
                elif ranged:
                    logger.info(f"Fetching pages {page_from}-{page_to}")
                    products, actual_start, actual_end, has_next = fetch_products_range(
                        store,
                        token,
                        page_from,
                        page_to,
                        on_progress=lambda m: (self._log(m), self._app.set_status(m)),
                    )
                    total = len(products)
                    if not products and actual_end < page_from:
                        raise ValueError(
                            f"Requested range {page_from}-{page_to} not found. Last available page is {actual_end}."
                        )
                    self._log(
                        f"Fetched pages {actual_start}-{actual_end} ({total} products). Next page exists: {'Yes' if has_next else 'No'}"
                    )
                    logger.info(f"Range download success: pages {actual_start}-{actual_end}, {total} products")
                else:
                    logger.info("Fetching all pages")
                    products = fetch_all_products(
                        store, token,
                        on_progress=lambda m: (self._log(m), self._app.set_status(m)),
                    )
                    total = len(products)
                    self._log(f"Fetched {total} products from all pages. Building Excel…")
                    logger.info(f"All pages download success: {total} products total")

                self._set_progress_step("fetch", "DONE", f"{total} products")

                # Enrich with metafields if any selected field requires them
                needs_mf = _needs_metafield_enrichment(fields)
                if needs_mf and products:
                    self._log(f"Fetching metafields for {total} products…")
                    self._set_progress_step("fetch", "RUNNING", "Fetching metafields")
                    enrich_products_with_metafields(
                        products, store, token, fields,
                        on_progress=lambda m: (self._log(m), self._app.set_status(m)),
                    )
                    self._set_progress_step("fetch", "DONE", f"{total} products + metafields")

                current_step = "images"
                self._set_progress_step("images", "SKIPPED", "Use Image Download tab")
                self._log("Image handling is managed in 'Image Download' tab.")

                self._app.start_prog("determinate", use_busy_cursor=False)

                logger.debug(f"Creating Excel workbook with {len(fields)} fields")
                current_step = "write"
                self._set_progress_step("write", "RUNNING", "Writing workbook")
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Products"
                for ci, f in enumerate(fields, 1):
                    cell = ws.cell(row=1, column=ci, value=f)
                    cell.font = Font(bold=True)

                BATCH = 500
                for start in range(0, total, BATCH):
                    chunk = products[start:start + BATCH]
                    for ri, p in enumerate(chunk, start + 2):
                        for ci, f in enumerate(fields, 1):
                            # Attempt primary path, then fallback between product-level
                            # and variant-level metafield locations when values are empty.
                            try:
                                raw_val = get_by_path(p, f)
                            except Exception:
                                raw_val = ""
                            val = raw_val if raw_val not in (None, "") else ""

                            if val == "" and isinstance(f, str) and f.startswith("metafields."):
                                # Try variant-level metafield (first variant)
                                alt = f"variants.0.{f}"
                                try:
                                    alt_val = get_by_path(p, alt)
                                    if alt_val not in (None, ""):
                                        val = alt_val
                                except Exception:
                                    pass

                            if val == "" and isinstance(f, str) and f.startswith("variants.0.metafields."):
                                # Try product-level metafield as fallback
                                alt = f[len("variants.0."):]
                                try:
                                    alt_val = get_by_path(p, alt)
                                    if alt_val not in (None, ""):
                                        val = alt_val
                                except Exception:
                                    pass

                            ws.cell(row=ri, column=ci, value=str(val))
                    self._app.set_prog_value(min(start + BATCH, total), total)
                    self._app.set_status(
                        f"Writing rows {start + 2}–{min(start + BATCH + 1, total + 1)}…"
                    )
                    self._set_progress_step("write", "RUNNING", f"Rows {start + 2}-{min(start + BATCH + 1, total + 1)}")

                wb.save(out)
                wb.close()
                logger.info(f"Excel file saved successfully: {out}")
                self._log(f"Saved: {out.resolve()}")

                self._set_progress_step("write", "DONE", f"{total} rows")
                current_step = "refresh"
                self._set_progress_step("refresh", "RUNNING", "Refreshing grid")
                self._app.refresh_live_database()
                self._set_progress_step("refresh", "DONE", "Live database refreshed")

                current_step = "backup"
                self._set_progress_step("backup", "RUNNING", "Creating backup file")
                bk = create_backup(out)
                logger.info(f"Backup created: {bk.name}")
                self._log(f"Backup: {bk.name}")
                self._app.notify_backup_created()
                self._set_progress_step("backup", "DONE", bk.name)

                ts = datetime.now().strftime("%Y-%m-%d  %H:%M:%S")
                self._last_dl.set(ts)
                if single:
                    self._log(f"Done! page {page_no}, {total} products, {len(fields)} fields.")
                    self._app.set_status(f"Download complete — page {page_no} ({total} products)")
                elif ranged:
                    self._log(f"Done! pages {page_from}-{page_to}, {total} products, {len(fields)} fields.")
                    self._app.set_status(f"Download complete — pages {page_from}-{page_to} ({total} products)")
                else:
                    self._log(f"Done! {total} products, {len(fields)} fields.")
                    self._app.set_status(f"Download complete — {total} products")

                self._set_progress_step("finish", "DONE", f"{total} products, {len(fields)} fields")
                self._set_progress_step("run", "DONE", "Completed")
                
                logger.info(f"Download completed successfully: {total} products, {len(fields)} fields")
                self._log(
                    "Download finished. Auto-open skipped to avoid Excel file lock; "
                    "use 'Open Current File' when needed."
                )

            except Exception as exc:
                logger.exception(f"Download error: {exc}")
                self._set_progress_step(current_step, "FAILED", str(exc))
                self._set_progress_step("finish", "FAILED", "Download failed")
                self._set_progress_step("run", "FAILED", "Stopped by error")
                self._log(f"ERROR: {exc}")
                self._log(
                    "Download FAIL context -> "
                    + json.dumps(
                        {
                            "mode": mode,
                            "page": page_no,
                            "from_page": page_from,
                            "to_page": page_to,
                            "selected_fields": fields,
                            "current_step": current_step,
                        },
                        ensure_ascii=False,
                        default=str,
                    )
                )
                self._app.set_status("Download failed")
                messagebox.showerror("Download Error", str(exc), parent=self)
            finally:
                self._running = False
                self._app.stop_prog(clear_busy_cursor=False)

        threading.Thread(target=task, daemon=True).start()

    def _open_current(self) -> None:
        p = self.current_xlsx()
        if not open_file(p):
            messagebox.showwarning(
                "Not Found", f"File not found:\n{p.resolve()}", parent=self
            )
